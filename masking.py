import argparse
import hashlib
import json
import os
import re

import openpyxl

try:
    import pymupdf as fitz  # PyMuPDF >= 1.24 preferred import name
except ImportError:
    import fitz

from presidio_analyzer import AnalyzerEngine, Pattern, PatternRecognizer
from presidio_anonymizer import AnonymizerEngine
from presidio_anonymizer.entities import OperatorConfig


# ==========================================
# 1. CORE PIPELINE ARCHITECTURE & VAULT
# ==========================================
class FinancialPrivacyEngine:
    # Matches tokens produced by generate_token, e.g. TOK_US_SSN_8B584CCF
    TOKEN_RE = re.compile(r"TOK_[A-Z0-9_]+_[0-9A-F]{8,}")

    # Bare numbers carry no linguistic signal, so NER entities (PERSON,
    # DATE_TIME, LOCATION, ...) are meaningless for numeric cells. Only
    # pattern/checksum-validated identifier types apply.
    NUMERIC_CELL_ENTITIES = [
        "CREDIT_CARD", "US_BANK_NUMBER", "US_SSN", "PHONE_NUMBER",
        "IBAN_CODE", "US_ITIN", "FINANCIAL_ACCOUNT",
    ]

    # Column rules: when a header cell matches, every cell below it in that
    # column is masked wholesale — no per-cell detection, full recall.
    # Order matters (first match wins): specific identifier types before
    # the broad name/address rules. Each rule is (entity, match, deny).
    COLUMN_RULES = [
        ("EMAIL_ADDRESS", re.compile(r"e-?mail", re.I), None),
        ("IP_ADDRESS", re.compile(r"\bip\b.{0,4}addr", re.I), None),
        ("PHONE_NUMBER",
         re.compile(r"phone|mobile|landline|telephone|fixed line|\bcell\b"
                    r"|\bfax\b|contact (no|number)", re.I), None),
        ("LK_NIC", re.compile(r"\bnic\b|national id", re.I), None),
        ("US_SSN", re.compile(r"\bssn\b|social security|tax id", re.I), None),
        ("LK_PASSPORT", re.compile(r"passport", re.I), None),
        ("FINANCIAL_ACCOUNT",
         re.compile(r"account\s*(no|number|#)|\biban\b|routing", re.I), None),
        ("DATE_TIME", re.compile(r"date of birth|\bdob\b|birth ?date", re.I), None),
        ("PERSON", re.compile(r"name\b|surname", re.I),
         re.compile(r"department|company|bank|branch|product|project|item"
                    r"|sheet|file|host|business|org|test|scheme|drug|road"
                    r"|street|city|country|holiday|brand|model", re.I)),
        ("LOCATION",
         re.compile(r"address|residence|home ?town|p\.?o\.? ?box|pobox", re.I), None),
    ]
    # Presidio ships many country-specific recognizers whose checksums fire
    # on arbitrary numbers (a 10-digit order id passes the UK NHS check ~10%
    # of the time; US driver-licence patterns match almost anything). None
    # are relevant to this deployment, so they are removed at startup.
    DISABLED_ENTITIES = {
        "UK_NHS", "UK_NINO", "AU_ABN", "AU_ACN", "AU_TFN", "AU_MEDICARE",
        "SG_NRIC_FIN", "SG_UEN", "IN_PAN", "IN_AADHAAR", "IN_VEHICLE_REGISTRATION",
        "IN_VOTER", "IN_PASSPORT", "ES_NIF", "ES_NIE", "IT_FISCAL_CODE",
        "IT_DRIVER_LICENSE", "IT_VAT_CODE", "IT_PASSPORT", "IT_IDENTITY_CARD",
        "PL_PESEL", "FI_PERSONAL_IDENTITY_CODE", "US_DRIVER_LICENSE",
        "MEDICAL_LICENSE", "KR_RRN", "TH_TNIN",
    }

    # Values that mean "no data" — never worth tokenizing in a rule column.
    NULL_MARKERS = {"", "-", "--", "n/a", "na", "none", "null", "nil"}
    # Shortest digit count any NUMERIC_CELL_ENTITIES pattern can match.
    MIN_NUMERIC_DIGITS = 8

    # Value-profile rules: if >= PROFILE_MIN_RATIO of a column's values match
    # ONE of these high-precision patterns (with validator), the column is an
    # identifier column no matter what its header says — or whether it has
    # one. Covers cryptic headers ("C3"), non-English headers, and headers
    # too deep to be found. Only formats distinctive enough to be safe.
    PROFILE_MIN_RATIO = 0.9
    PROFILE_MIN_SAMPLES = 4

    @staticmethod
    def _valid_lk_nic(v):
        # NIC embeds day-of-year (001-366, +500 for women) after the year.
        digits = v[:-1] if v[-1] in "VvXx" else v
        day = int(digits[2:5]) if len(digits) == 9 else int(digits[4:7])
        return 1 <= day <= 366 or 501 <= day <= 866

    @staticmethod
    def _valid_luhn(v):
        ds = [int(c) for c in v if c.isdigit()]
        if not 13 <= len(ds) <= 19:
            return False
        total = 0
        for i, d in enumerate(reversed(ds)):
            if i % 2 == 1:
                d *= 2
                if d > 9:
                    d -= 9
            total += d
        return total % 10 == 0

    PROFILE_PATTERNS = [
        ("LK_NIC", re.compile(r"(?:\d{9}[VvXx]|(?:19|20)\d{10})"), "_valid_lk_nic"),
        ("EMAIL_ADDRESS", re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]{2,})+"), None),
        ("PHONE_NUMBER",
         re.compile(r"(?:\+94[- ]?|0)[1-9]\d[- ]?\d{3}[- ]?\d{4}|\+\d{1,3}[- ]?\d{2,4}[- ]?\d{3}[- ]?\d{3,4}"),
         None),
        ("CREDIT_CARD", re.compile(r"\d[\d -]{11,21}\d"), "_valid_luhn"),
        ("US_SSN", re.compile(r"\d{3}-\d{2}-\d{4}"), "_valid_ssn"),
        ("LK_PASSPORT", re.compile(r"[NDS]\d{7}"), None),
    ]

    @staticmethod
    def _valid_ssn(v):
        area, group, serial = v.split("-")
        return (area not in ("000", "666") and not area.startswith("9")
                and group != "00" and serial != "0000")

    def _profile_columns(self, ws, skip_cols, header_row, end_row):
        """{column: (entity, header_row_or_0)} for columns whose values
        (between header_row and end_row) overwhelmingly match one
        identifier pattern."""
        found = {}
        start = (header_row or 0) + 1
        if end_row < start:
            return found
        for col_cells in ws.iter_cols(min_row=start, max_row=end_row):
            col = col_cells[0].column
            if col in skip_cols:
                continue
            values = []
            for c in col_cells:
                v = c.value
                if v is None or isinstance(v, bool):
                    continue
                if isinstance(v, (int, float)):
                    v = str(int(v)) if float(v).is_integer() else str(v)
                v = str(v).strip()
                if v and v.lower() not in self.NULL_MARKERS and not self.TOKEN_RE.fullmatch(v):
                    values.append(v)
            if len(values) < self.PROFILE_MIN_SAMPLES:
                continue
            for entity, pat, validator in self.PROFILE_PATTERNS:
                check = getattr(self, validator) if validator else (lambda v: True)
                hits = sum(1 for v in values if pat.fullmatch(v) and check(v))
                if hits / len(values) >= self.PROFILE_MIN_RATIO:
                    found[col] = (entity, header_row or 0)
                    break
        return found

    # Words that signal a phone number nearby. Presidio's default list lacks
    # landline vocabulary and messaging cues (OTP/SMS imply a mobile).
    # NOTE: Presidio matches context words as SUBSTRINGS of the context
    # tokens ("ring" would match "string", "line" matches "online"), so
    # only words too long/specific to hide inside other words belong here.
    PHONE_CONTEXT = ["phone", "number", "telephone", "cellphone", "mobile",
                     "call", "landline", "fixed", "dial", "contact", "fax",
                     "hotline", "otp", "sms", "whatsapp", "tel"]

    @staticmethod
    def _is_headerish(value):
        if not isinstance(value, str):
            return False
        t = value.strip()
        return bool(t) and len(t) <= 40 and len(t.split()) <= 6

    def _find_header_row(self, ws, scan_rows=40):
        """Row number of the first top row holding >=2 short header-like
        strings, or None."""
        for row in ws.iter_rows(min_row=1, max_row=min(scan_rows, ws.max_row)):
            if (sum(1 for c in row if self._is_headerish(c.value)) >= 2
                    and not self._looks_like_data(row)):
                return row[0].row
        return None

    def _looks_like_data(self, row):
        """A row holding numbers or identifier-shaped values is data, not a
        header — guards the stacked-table detection below."""
        for c in row:
            v = c.value
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                return True
            if isinstance(v, str):
                t = v.strip()
                if any(p.fullmatch(t) for _, p, _ in self.PROFILE_PATTERNS):
                    return True
        return False

    def _sheet_segments(self, ws):
        """
        Split a sheet into tables. The first header row is found in the top
        rows; a later row that is header-like, follows an empty row, and
        holds no data-shaped values starts a new table (stacked tables on
        one sheet). Each segment carries its own column rules and context.
        Returns [{header, end, rules, context}] ordered by row.
        """
        first = self._find_header_row(ws)
        headers = []
        if first is not None:
            headers.append(first)
            prev_empty = False
            for row in ws.iter_rows(min_row=first + 1):
                nonempty = [c for c in row if c.value is not None and str(c.value).strip()]
                if not nonempty:
                    prev_empty = True
                    continue
                if (prev_empty and sum(1 for c in row if self._is_headerish(c.value)) >= 2
                        and not self._looks_like_data(row)):
                    headers.append(row[0].row)
                prev_empty = False
        if not headers:
            return [{"header": 0, "end": ws.max_row,
                     "rules": {c: (r[0], "value-profile") for c, r in
                               self._profile_columns(ws, set(), None, ws.max_row).items()},
                     "context": self._column_context(ws, None)}]
        segments = []
        for i, h in enumerate(headers):
            end = headers[i + 1] - 1 if i + 1 < len(headers) else ws.max_row
            rules = {}
            for cell in ws[h]:
                if not self._is_headerish(cell.value):
                    continue
                t = cell.value.strip()
                for entity, pat, deny in self.COLUMN_RULES:
                    if pat.search(t) and not (deny and deny.search(t)):
                        if entity == "DATE_TIME" and self.dates == "none":
                            break
                        rules[cell.column] = (entity, "column-rule")
                        break
            for col, r in self._profile_columns(ws, set(rules), h, end).items():
                rules[col] = (r[0], "value-profile")
            segments.append({"header": h, "end": end, "rules": rules,
                             "context": self._column_context(ws, h)})
        return segments

    def __init__(self, salt=None, min_score=0.6, entities=None, language="en",
                 nlp_model=None, dates="birth", column_rules=True):
        """
        salt:      secret used for deterministic tokens. Prefer the
                   PII_TOKEN_SALT environment variable over hardcoding.
        min_score: minimum analyzer confidence required to mask a match.
        entities:  optional list of entity types to detect (None = all).
        nlp_model: spaCy model for NER (None = Presidio's default,
                   en_core_web_lg). "en_core_web_trf" gives noticeably
                   better name recall at a higher runtime cost.
        dates:     date-masking policy. "birth" (default) masks only dates
                   in a birth context (a DOB is a strong quasi-identifier;
                   ordinary transaction/event dates stay analyzable),
                   "all" masks every detected date (HIPAA-style), "none"
                   masks no dates.
        """
        if dates not in ("birth", "all", "none"):
            raise ValueError('dates must be "birth", "all", or "none"')
        self.dates = dates
        self.column_rules = column_rules
        self.salt = salt or os.environ.get("PII_TOKEN_SALT", "EnterpriseRiskManagement2026")
        self.min_score = min_score
        self.entities = entities
        self.language = language
        self.vault = {}  # {token: original_value} — treat as sensitive material
        self.numeric_tokens = set()  # tokens whose source cell was numeric (fallback)
        self.numeric_cells = set()   # "Sheet!A1" cells that were numeric (authoritative)
        self.report = []  # per-run detection details: where/entity/score/text
        self._last_findings = []  # findings of the most recent analyze call

        if nlp_model:
            from presidio_analyzer.nlp_engine import NlpEngineProvider
            provider = NlpEngineProvider(nlp_configuration={
                "nlp_engine_name": "spacy",
                "models": [{"lang_code": language, "model_name": nlp_model}],
            })
            self.analyzer = AnalyzerEngine(nlp_engine=provider.create_engine())
        else:
            self.analyzer = AnalyzerEngine()
        self.anonymizer = AnonymizerEngine()
        for rec in list(self.analyzer.registry.recognizers):
            if set(rec.supported_entities) <= self.DISABLED_ENTITIES:
                self.analyzer.registry.remove_recognizer(rec.name)
        self.setup_custom_financial_matchers()
        self.setup_sri_lanka_recognizers()

    def setup_custom_financial_matchers(self):
        """Register account-number recognizers with context-aware scoring."""
        patterns = [
            # Structured ids like CHQ-4491-002: distinctive shape, high base score.
            Pattern(
                name="structured_account_pattern",
                regex=r"\b[A-Z]{2,4}-\d{3,6}-\d{2,4}\b",
                score=0.65,
            ),
            # Bare long digit runs are ambiguous (order ids, timestamps, ...).
            # Low base score: only masked when nearby context words boost it
            # past min_score via Presidio's context enhancer.
            Pattern(
                name="bare_account_number_pattern",
                regex=r"\b\d{9,18}\b",
                score=0.3,
            ),
            # Short labelled refs like EPF/ETF numbers (A/12345, E-88123).
            # Base score below threshold; only masked when column/sentence
            # context (epf, fund, account, ...) boosts it.
            Pattern(
                name="labelled_ref_pattern",
                regex=r"\b[A-Z]{1,2}[/-]\d{4,6}\b",
                score=0.3,
            ),
        ]
        acct_recognizer = PatternRecognizer(
            supported_entity="FINANCIAL_ACCOUNT",
            patterns=patterns,
            context=["account", "acct", "iban", "routing", "bsb", "swift",
                     "ledger", "policy", "member", "customer",
                     "epf", "etf", "provident", "fund"],
        )
        self.analyzer.registry.add_recognizer(acct_recognizer)

        # OCR-tolerant email fallback: the built-in recognizer validates the
        # TLD, so an OCR misread like ".Ik" for ".lk" makes it reject the
        # whole address. For masking, over-matching beats leaking.
        loose_email = PatternRecognizer(
            supported_entity="EMAIL_ADDRESS",
            patterns=[Pattern("email_loose",
                              r"\b[\w.+-]+@[\w-]+(?:\.[\w-]{2,})+\b", 0.7)],
        )
        self.analyzer.registry.add_recognizer(loose_email)

        # A capitalized word right after an honorific or role noun is a
        # strong person signal even when NER does not know the name
        # ("overseer Jayasuriya", "Mr. Silva"). Catches names statistical
        # NER misses, which hits non-Western names hardest.
        titled_name = PatternRecognizer(
            supported_entity="PERSON",
            patterns=[Pattern(
                "role_titled_name",
                r"(?i:\b(?:mr|mrs|ms|miss|dr|hon|rev|prof|overseer|inspector|"
                r"officer|constable|sergeant|clerk|agent|foreman|headman|"
                r"mudaliyar|arachchi|manager|supervisor|engineer|surveyor)"
                r"\b\.?\s+)(?-i:[A-Z][\w'’-]{2,})",
                0.6)],
        )
        self.analyzer.registry.add_recognizer(titled_name)

        # Inverted all-caps names ("JAYAWARDENA, SANDUNI D") — the payroll /
        # HR export format that defeats statistical NER. Requires the comma
        # and an optional trailing initial, so ordinary caps headings don't
        # match.
        inverted_caps_name = PatternRecognizer(
            supported_entity="PERSON",
            patterns=[Pattern(
                "inverted_caps_name",
                r"\b[A-Z][A-Z'’-]{3,}(?: [A-Z][A-Z'’-]{2,})?, [A-Z][A-Z'’-]{2,}"
                r"(?: [A-Z][A-Z'’-]{2,})?(?: [A-Z]\.?)?(?=[\s,.;:)]|$)(?!, [A-Z])",
                0.6)],
            # Presidio compiles patterns with IGNORECASE by default, which
            # would turn this caps-only heuristic into "any Word, Word".
            global_regex_flags=re.DOTALL | re.MULTILINE,
        )
        self.analyzer.registry.add_recognizer(inverted_caps_name)

    def setup_sri_lanka_recognizers(self):
        """Recognizers for Sri Lankan identifier formats."""
        nic_recognizer = PatternRecognizer(
            supported_entity="LK_NIC",
            patterns=[
                # Old NIC: 9 digits + V/X suffix — distinctive on its own.
                Pattern("lk_nic_old", r"\b\d{9}[VvXx]\b", 0.8),
                # New NIC: 12 digits starting with the birth year. Ambiguous
                # with other long numbers, so it needs context to pass.
                Pattern("lk_nic_new", r"\b(?:19|20)\d{10}\b", 0.45),
            ],
            context=["nic", "national", "identity"],
        )
        passport_recognizer = PatternRecognizer(
            supported_entity="LK_PASSPORT",
            # N/D/S prefix + 7 digits; too generic alone, requires context.
            patterns=[Pattern("lk_passport", r"\b[NDS]\d{7}\b", 0.3)],
            context=["passport", "travel"],
        )
        # Sri Lankan phone formats. A formatted MOBILE (07X / +94 7X) is a
        # personal identifier distinctive enough to mask without context;
        # landlines are often business numbers, so they stay below the
        # threshold until context (a header, 'call', 'tel'...) lifts them.
        lk_phone_recognizer = PatternRecognizer(
            supported_entity="PHONE_NUMBER",
            patterns=[
                Pattern("lk_mobile",
                        r"(?<!\d)(?:\+94[- ]?|0)7\d[- ]?\d{3}[- ]?\d{4}(?!\d)", 0.65),
                Pattern("lk_landline",
                        r"(?<!\d)(?:\+94[- ]?|0)[1-9]\d[- ]?\d{3}[- ]?\d{4}(?!\d)", 0.45),
            ],
            context=self.PHONE_CONTEXT,
        )
        self.analyzer.registry.add_recognizer(nic_recognizer)
        self.analyzer.registry.add_recognizer(passport_recognizer)
        self.analyzer.registry.add_recognizer(lk_phone_recognizer)

        # Re-register the phone recognizer with Sri Lanka included, so the
        # phonenumbers library validates +94 / 0XX-XXXXXXX formats.
        from presidio_analyzer.predefined_recognizers import PhoneRecognizer
        self.analyzer.registry.remove_recognizer("PhoneRecognizer")
        self.analyzer.registry.add_recognizer(
            PhoneRecognizer(
                supported_regions=("LK", "US", "GB", "IN"),
                context=self.PHONE_CONTEXT,
            )
        )

    def generate_token(self, original_text, entity_type):
        """Deterministic, repeatable token for a value (same input -> same token)."""
        clean_text = str(original_text).strip()
        digest = hashlib.sha256((clean_text + self.salt).encode("utf-8")).hexdigest().upper()

        # Extend the hash prefix on the (rare) collision with a different value.
        length = 8
        token = f"TOK_{entity_type}_{digest[:length]}"
        while token in self.vault and self.vault[token] != clean_text:
            length += 4
            token = f"TOK_{entity_type}_{digest[:length]}"

        self.vault[token] = clean_text
        return token

    def analyze_text(self, text, entities=None, context=None):
        """
        Run the analyzer and filter out known-spurious matches. Shared by
        the Excel and PDF paths so both apply the same rules.
        """
        self.analyzer_calls = getattr(self, "analyzer_calls", 0) + 1
        results = self.analyzer.analyze(
            text=text,
            language=self.language,
            entities=entities if entities is not None else self.entities,
            context=context,
            score_threshold=self.min_score,
        )
        # NER labels (person/place/group) on text with no letters are
        # model noise — the transformer model in particular will happily
        # tag a bare salary figure as PERSON. Pattern-based entities
        # (accounts, cards, ...) are unaffected.
        results = [
            r for r in results
            if r.entity_type not in ("PERSON", "LOCATION", "NRP")
            or any(c.isalpha() for c in text[r.start:r.end])
        ]

        # Names never contain digits: a PERSON span with a digit is a code
        # (EMP-100, WP CAB-1234, PRD-5521) that NER mistook for a name. And
        # a lone token of <=3 letters ("Max", "Pro", "Jan") standing as a
        # whole value is an abbreviation far more often than a person.
        def plausible_person(span):
            if any(c.isdigit() for c in span):
                return False
            words = span.split()
            return not (len(words) == 1 and len(words[0].strip(".,")) <= 3
                        and span.strip() == text.strip())
        results = [
            r for r in results
            if (r.entity_type != "PERSON" or plausible_person(text[r.start:r.end]))
            and not (r.entity_type == "NRP" and any(c.isdigit() for c in text[r.start:r.end]))
        ]

        # A "person" whose final word is a role/honorific ("Hon. Attorney")
        # is a title fragment, not a name — drop it.
        results = [
            r for r in results
            if not (
                r.entity_type == "PERSON"
                and text[r.start:r.end].split()
                and text[r.start:r.end].split()[-1].strip(".,;:'\"()").lower()
                in self._NAME_PROPAGATION_STOPWORDS
            )
        ]

        # Date policy. Dates are analytical fields, not direct identifiers,
        # so by default only birth-linked dates are masked (a DOB is a
        # classic re-identification quasi-identifier); every other date is
        # left for analysis. "all" masks every date, "none" masks none.
        # Digit-only runs (1919, 71829, a salary of 165624) are never
        # accepted as dates unless a birth context says so AND they have a
        # date-like length — NER routinely mislabels bare numbers, and
        # dropping the date label lets the real recognizer (e.g.
        # FINANCIAL_ACCOUNT) win instead.
        birth_re = re.compile(r"\b(dob|birth|born|birthday)\b", re.IGNORECASE)
        ctx_birth = bool(context) and any(birth_re.search(w) for w in context)

        def keep_date(r):
            if self.dates == "none":
                return False
            snippet = text[r.start:r.end].strip()
            if snippet.isdigit():
                return ctx_birth and len(snippet) in (6, 8)
            if self.dates == "all":
                return True
            window = text[max(0, r.start - 60):min(len(text), r.end + 30)]
            return ctx_birth or bool(birth_re.search(window))

        return [r for r in results
                if r.entity_type != "DATE_TIME" or keep_date(r)]

    def pseudonymize_text(self, text, entities=None, context=None):
        """
        Replace only the PII spans inside `text` with tokens, keeping the
        surrounding content intact. Returns (new_text, changed).

        context: extra words (e.g. a spreadsheet column header) fed to
        Presidio's context enhancer, which raises the score of ambiguous
        matches such as bare phone or account numbers.
        """
        results = self.analyze_text(text, entities=entities, context=context)
        self._last_findings = [
            {"entity": r.entity_type, "score": round(r.score, 2),
             "text": text[r.start:r.end]}
            for r in results
        ]

        # Ignore any span overlapping an existing token (makes re-runs
        # idempotent: tokens and their surroundings are never re-masked).
        token_spans = [m.span() for m in self.TOKEN_RE.finditer(text)]
        results = [
            r for r in results
            if not any(r.start < end and r.end > start for start, end in token_spans)
        ]
        if not results:
            return text, False

        # One operator per detected entity type; the anonymizer resolves
        # overlapping matches before applying them.
        operators = {
            entity_type: OperatorConfig(
                "custom",
                {"lambda": lambda t, et=entity_type: self.generate_token(t, et)},
            )
            for entity_type in {r.entity_type for r in results}
        }
        outcome = self.anonymizer.anonymize(
            text=text, analyzer_results=results, operators=operators
        )
        return outcome.text, outcome.text != text

    def depseudonymize_text(self, text):
        """Restore original values for any vault tokens present in `text`."""
        return self.TOKEN_RE.sub(lambda m: self.vault.get(m.group(0), m.group(0)), text)

    # -------- Vault persistence (plaintext JSON: store securely) --------
    def save_vault(self, path):
        payload = {"mappings": self.vault, "numeric_tokens": sorted(self.numeric_tokens),
                   "numeric_cells": sorted(self.numeric_cells)}
        with open(path, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, indent=2)
        print(f"[Vault] {len(self.vault)} mappings written to: {path} (protect this file)")

    def load_vault(self, path):
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        if "mappings" in data:
            self.vault.update(data["mappings"])
            self.numeric_tokens.update(data.get("numeric_tokens", []))
            self.numeric_cells.update(data.get("numeric_cells", []))
        else:  # legacy flat {token: value} format
            self.vault.update(data)

# ==========================================
# 2. EXCEL MULTI-SHEET ENGINE (openpyxl)
# ==========================================
    def _column_context(self, ws, header_row, max_rows=8, max_words=10):
        """
        Per-column context words for the analyzer, so a bare phone number in
        a column titled 'Phone Number' scores as it would inside a sentence.

        Context must come from labels, never from data: when a header row is
        detected, each column's context is its header cell plus the cell
        directly above it (two-row headers like 'Contact' / 'Mobile'). Only
        when no header row can be found do we fall back to harvesting the
        top rows wholesale.
        """
        words_of = lambda v: [w.lower() for w in re.findall(r"[A-Za-z]{3,}", v)] \
            if isinstance(v, str) else []
        context = {}
        if header_row is not None:
            for cell in ws[header_row]:
                words = words_of(cell.value)
                if header_row > 1:
                    words = words_of(ws.cell(header_row - 1, cell.column).value) + words
                if words:
                    context[cell.column] = words[:max_words]
            return context

        for col_cells in ws.iter_cols(max_row=min(max_rows, ws.max_row)):
            words = []
            for cell in col_cells:
                words.extend(words_of(cell.value))
            if words:
                context[col_cells[0].column] = words[:max_words]
        return context

    def pseudonymize_excel(self, input_path, output_path):
        """Scan every sheet and cell, replacing detected PII spans with tokens."""
        wb = openpyxl.load_workbook(input_path)
        cells_changed = 0

        numeric_entities = self.NUMERIC_CELL_ENTITIES
        if self.entities is not None:
            numeric_entities = [e for e in numeric_entities if e in self.entities]

        for ws in wb.worksheets:
            segments = self._sheet_segments(ws)
            if not self.column_rules:
                for sg in segments:
                    sg["rules"] = {}
            cache = {}  # (text, numeric?, column, segment) -> (new_text, changed, findings)
            seg_iter = iter(segments)
            sg = next(seg_iter, None)
            for row in ws.iter_rows():
                for cell in row:
                    # advance to the segment containing this row (rows above
                    # the first header belong to no segment)
                    while sg is not None and cell.row > sg["end"]:
                        sg = next(seg_iter, None)
                    in_seg = sg is not None and cell.row > sg["header"]
                    if sg is not None and cell.row == sg["header"]:
                        continue  # header cells are labels, never data
                    value = cell.value
                    entities = None
                    numeric = isinstance(value, (int, float)) and not isinstance(value, bool)

                    # Column rule: mask the whole column below its header —
                    # any value type, no detection needed.
                    rule = sg["rules"].get(cell.column) if in_seg else None
                    if rule and value is not None:
                        rule_text = (str(int(value)) if numeric and float(value).is_integer()
                                     else str(value)).strip()
                        if (rule_text.lower() in self.NULL_MARKERS
                                or self.TOKEN_RE.fullmatch(rule_text)):
                            continue
                        token = self.generate_token(rule_text, rule[0])
                        if numeric:
                            self.numeric_tokens.add(token)
                            self.numeric_cells.add(f"{ws.title}!{cell.coordinate}")
                        cell.value = token
                        cells_changed += 1
                        self.report.append(
                            {"where": f"{ws.title}!{cell.coordinate}",
                             "method": rule[1], "entity": rule[0],
                             "score": None, "text": rule_text})
                        continue

                    if isinstance(value, str):
                        text = value
                        # Pre-filter: nothing shorter than 3 chars or without
                        # any letter/digit can be PII — skip the analyzer.
                        stripped = text.strip()
                        if len(stripped) < 3 or not any(c.isalnum() for c in stripped):
                            continue
                    elif numeric:
                        # Card/account numbers stored as numeric cells must
                        # still be analyzed, but only against identifier
                        # patterns — ordinary figures are left untouched.
                        text = str(int(value)) if float(value).is_integer() else str(value)
                        entities = numeric_entities
                        # Pre-filter: every numeric identifier type needs at
                        # least 8 digits (bank 8+, SSN 9, NIC 12, cards 13+),
                        # so shorter numbers (amounts, counts, years) skip
                        # the analyzer entirely.
                        if sum(c.isdigit() for c in text) < self.MIN_NUMERIC_DIGITS:
                            continue
                    else:
                        continue

                    # Per-run cache: spreadsheets repeat values (cities,
                    # departments, statuses); each distinct (value, entity
                    # set, column context) is analyzed once. Results are
                    # identical by construction — same input, same output.
                    context = sg["context"].get(cell.column) if in_seg else None
                    key = (text, entities is not None,
                           (cell.column, sg["header"]) if context else None)
                    hit = cache.get(key)
                    if hit is None:
                        new_text, changed = self.pseudonymize_text(
                            text, entities=entities, context=context
                        )
                        cache[key] = (new_text, changed, list(self._last_findings))
                    else:
                        new_text, changed, self._last_findings = hit
                    if changed:
                        if entities is not None and self.TOKEN_RE.fullmatch(new_text):
                            self.numeric_tokens.add(new_text)
                            self.numeric_cells.add(f"{ws.title}!{cell.coordinate}")
                        cell.value = new_text
                        cells_changed += 1
                        for f in self._last_findings:
                            self.report.append(
                                {"where": f"{ws.title}!{cell.coordinate}",
                                 "method": "detected", **f})

        wb.save(output_path)
        print(f"[Success] Excel saved to: {output_path} ({cells_changed} cells modified)")

    def depseudonymize_excel(self, input_path, output_path):
        """
        Reverse a masked workbook using the loaded vault. Cells recorded as
        numeric at masking time are restored to numbers; everything else
        stays text, so identifiers like NICs keep their original type.
        """
        wb = openpyxl.load_workbook(input_path)
        cells_restored = 0
        unresolved = 0

        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if not isinstance(cell.value, str) or not self.TOKEN_RE.search(cell.value):
                        continue
                    # Per-cell record wins (the same value can be a number in
                    # one sheet and text in another); per-token is the
                    # fallback when the vault came from a different workbook.
                    if self.numeric_cells:
                        was_numeric = f"{ws.title}!{cell.coordinate}" in self.numeric_cells
                    else:
                        was_numeric = cell.value.strip() in self.numeric_tokens
                    restored = self.depseudonymize_text(cell.value)
                    if self.TOKEN_RE.search(restored):
                        unresolved += 1  # token missing from the vault
                    if restored != cell.value:
                        if was_numeric:
                            cell.value = float(restored) if "." in restored else int(restored)
                        else:
                            cell.value = restored
                        cells_restored += 1

        wb.save(output_path)
        print(f"[Success] Restored workbook saved to: {output_path} "
              f"({cells_restored} cells restored)")
        if unresolved:
            print(f"[Warning] {unresolved} cells still contain tokens not found "
                  f"in the vault — was the correct vault loaded?")

# ==========================================
# 3. PDF SPATIAL MASKING ENGINE (PyMuPDF)
# ==========================================
    @staticmethod
    def _norm_token(token):
        """Normalize a word for matching: strip edge punctuation and a
        possessive suffix, casefold."""
        token = token.strip("\"'’‘`.,;:()[]{}<>!?—–-*")
        token = re.sub(r"['’‘`]s$", "", token)
        return token.casefold()

    # Words never worth redacting on their own when propagating name parts:
    # honorifics and legal/role vocabulary NER sometimes folds into a PERSON.
    _NAME_PROPAGATION_STOPWORDS = {
        "hon", "mr", "mrs", "ms", "dr", "prof", "rev", "attorney", "general",
        "justice", "court", "judge", "counsel", "learned", "president",
        "secretary", "chairman", "director", "officer", "petitioner",
        "respondent", "appellant", "accused", "complainant",
    }

    @staticmethod
    def _find_tessdata():
        """Locate the Tesseract language-data directory for OCR."""
        prefix = os.environ.get("TESSDATA_PREFIX")
        if prefix and os.path.isdir(prefix):
            return prefix
        import glob
        hits = glob.glob("/usr/share/tesseract-ocr/*/tessdata") \
            + glob.glob("/usr/local/share/tessdata") \
            + glob.glob("/usr/share/tessdata")
        return hits[0] if hits else None

    def _textpage_for(self, page, ocr_dpi=300):
        """
        Return an OCR textpage for scanned pages (image content but no
        usable text layer), or None when the native text layer suffices.
        """
        if len(page.get_text("text").strip()) >= 30:
            return None
        if not page.get_images(full=True):
            return None
        tessdata = self._find_tessdata()
        if tessdata is None:
            print(f"[Warning] Page {page.number + 1} looks scanned but Tesseract "
                  f"language data was not found — install tesseract-ocr. "
                  f"This page will NOT be redacted.")
            return None
        print(f"[OCR] Page {page.number + 1} has no text layer — running OCR")
        return page.get_textpage_ocr(language="eng", dpi=ocr_dpi, full=True,
                                     tessdata=tessdata)

    def redact_spatial_pdf(self, input_path, output_path):
        """
        Redact a PDF in two passes. Pass 1 detects entities over each page's
        whitespace-normalized text (layouts that break names across lines
        defeat NER otherwise). Pass 2 blacks out every occurrence of every
        detected value on every page — so a name NER only caught once is
        still removed wherever else it appears — and scrubs the underlying
        character stream. Individual words of detected person names are
        propagated too, catching partial mentions of the same person.
        Scanned pages are OCR'd and the matching image pixels are destroyed,
        not merely covered.
        """
        doc = fitz.open(input_path)

        # Keep one Page object per page alive for the whole run: an OCR
        # textpage weak-references its page and dies with it otherwise.
        pages = [doc[i] for i in range(len(doc))]
        textpages = [self._textpage_for(page) for page in pages]
        page_texts = [
            re.sub(r"\s+", " ", page.get_text("text", textpage=tp))
            for page, tp in zip(pages, textpages)
        ]

        # Pass 1: document-wide detection.
        snippets = {}  # {text: entity_type}
        for page_no, text in enumerate(page_texts, 1):
            for result in self.analyze_text(text):
                snippet = text[result.start:result.end].strip(" .,;:'\"()")
                if len(snippet) >= 3:
                    snippets.setdefault(snippet, result.entity_type)
                    self.report.append(
                        {"where": f"page {page_no}", "method": "detected",
                         "entity": result.entity_type,
                         "score": round(result.score, 2), "text": snippet})

        # Propagate parts of person names (e.g. a surname on its own).
        name_words = set()
        for snippet, entity_type in list(snippets.items()):
            if entity_type != "PERSON":
                continue
            for word in snippet.split():
                word = word.strip(".,;:'\"()")
                if (len(word) >= 4 and word[0].isupper()
                        and word.lower() not in self._NAME_PROPAGATION_STOPWORDS):
                    snippets.setdefault(word, "PERSON")
                    name_words.add(word.lower())

        # Expand to capitalized word-runs containing a known name word, so
        # "Murshida Shiyam" is fully removed even if NER only ever saw
        # "Shiyam" as part of another person's name.
        run_re = re.compile(r"\b[A-Z][\w'’-]*(?:\s+[A-Z][\w'’-]*)+")
        for text in page_texts:
            for match in run_re.finditer(text):
                run = match.group(0)
                if any(w.strip(".,;:'\"()").lower() in name_words for w in run.split()):
                    snippets.setdefault(run, "PERSON")

        # Report snippets added by propagation/expansion (no analyzer score).
        directly_detected = {e["text"] for e in self.report if e["method"] == "detected"}
        for snippet, entity_type in snippets.items():
            if snippet not in directly_detected:
                self.report.append(
                    {"where": "document", "method": "propagated",
                     "entity": entity_type, "score": None, "text": snippet})

        # Pass 2: redact every occurrence of every detected value. Rects
        # come from matching the snippet's tokens against the page's word
        # sequence — geometry-only checks are unreliable on skewed scans
        # where OCR word boxes can span lines.
        redactions = 0
        for page, tp in zip(pages, textpages):
            word_boxes = [(fitz.Rect(w[:4]), w[4])
                          for w in page.get_text("words", textpage=tp)]
            norm_words = [self._norm_token(w) for _, w in word_boxes]

            def word_match_rects(snippet):
                tokens = [t for t in map(self._norm_token, snippet.split()) if t]
                if not tokens:
                    return []
                rects = []
                for i in range(len(norm_words) - len(tokens) + 1):
                    if norm_words[i:i + len(tokens)] == tokens:
                        rects.extend(r for r, _ in word_boxes[i:i + len(tokens)])
                return rects

            def add_redaction(rect):
                if tp is not None:
                    # OCR boxes sit tighter than the printed glyphs; pad a
                    # little so no character fringes survive.
                    rect = fitz.Rect(rect.x0 - 2, rect.y0 - 1,
                                     rect.x1 + 2, rect.y1 + 1)
                page.add_redact_annot(rect, fill=(0, 0, 0))

            for snippet, entity_type in snippets.items():
                rects = word_match_rects(snippet)
                # Fallback for values embedded inside a larger word (native
                # text) or merged OCR tokens: substring search, restricted
                # to substantial snippets so short false positives cannot
                # chew through unrelated words.
                if not rects and len(snippet) >= 5:
                    rects = page.search_for(snippet, textpage=tp)
                for rect in rects:
                    add_redaction(rect)
                    redactions += 1
                if rects:
                    # Record what was removed so an audit trail exists.
                    self.generate_token(snippet, entity_type)
            # PDF_REDACT_IMAGE_PIXELS erases matching pixels inside scanned
            # images, so the PII is destroyed rather than covered.
            page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_PIXELS)

        doc.save(output_path)
        print(f"[Success] PDF saved to: {output_path} ({redactions} regions redacted)")

# ==========================================
# 4. COMMAND-LINE ENTRY POINT
# ==========================================
def main():
    parser = argparse.ArgumentParser(
        description="Pseudonymize Excel workbooks / redact PDFs with Presidio."
    )
    parser.add_argument("input", help="Input .xlsx or .pdf file")
    parser.add_argument("output", nargs="?", help="Output path (default: <input>_masked.<ext>)")
    parser.add_argument("--vault", help="Vault JSON file: written after masking, "
                        "read when using --restore")
    parser.add_argument("--restore", action="store_true",
                        help="Reverse a masked .xlsx using the vault given via --vault")
    parser.add_argument("--min-score", type=float, default=0.6,
                        help="Minimum detection confidence (default: 0.6)")
    parser.add_argument("--entities", nargs="*", default=None,
                        help="Restrict detection to these entity types (default: all)")
    parser.add_argument("--nlp-model", default=None,
                        help="spaCy model for NER, e.g. en_core_web_trf "
                             "(default: Presidio's en_core_web_lg)")
    parser.add_argument("--dates", choices=["birth", "all", "none"], default="birth",
                        help="Date masking policy: 'birth' masks only birth-linked "
                             "dates (default), 'all' masks every date, 'none' keeps all")
    parser.add_argument("--no-column-rules", action="store_true",
                        help="Disable header-based column rules (mask whole 'Name'/"
                             "'Address'/... columns without per-cell detection)")
    args = parser.parse_args()

    root, ext = os.path.splitext(args.input)
    output = args.output or f"{root}_masked{ext}"

    engine = FinancialPrivacyEngine(min_score=args.min_score, entities=args.entities,
                                    nlp_model=args.nlp_model, dates=args.dates,
                                    column_rules=not args.no_column_rules)

    ext = ext.lower()
    if args.restore:
        if not args.vault:
            parser.error("--restore requires --vault <mapping file>")
        if ext not in (".xlsx", ".xlsm"):
            parser.error("--restore only supports Excel files (PDF redaction is permanent)")
        engine.load_vault(args.vault)
        engine.depseudonymize_excel(args.input, args.output or f"{root}_restored{ext}")
        return

    if ext in (".xlsx", ".xlsm"):
        engine.pseudonymize_excel(args.input, output)
    elif ext == ".pdf":
        engine.redact_spatial_pdf(args.input, output)
    else:
        parser.error(f"Unsupported file type: {ext} (expected .xlsx, .xlsm, or .pdf)")

    if args.vault:
        engine.save_vault(args.vault)


if __name__ == "__main__":
    main()
