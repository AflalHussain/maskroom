"""Excel pipeline (openpyxl): header detection, stacked-table segmentation,
column rules, value profiling, span-level pseudonymization and restore."""
import re

import openpyxl

from . import rules


class ExcelMixin:
    # ------------------------------------------------------ table layout
    @staticmethod
    def _is_headerish(value):
        if not isinstance(value, str):
            return False
        t = value.strip()
        return bool(t) and len(t) <= 40 and len(t.split()) <= 6

    def _find_header_row(self, ws, scan_rows=40):
        """Row number of the first top row holding >=2 short header-like
        strings, or None."""
        for row in ws.iter_rows(min_row=1, max_row=min(scan_rows, ws.max_row)):
            if (sum(1 for c in row if self._is_headerish(c.value)) >= 2
                    and not self._looks_like_data(row)):
                return row[0].row
        return None

    def _looks_like_data(self, row):
        """A row holding numbers or identifier-shaped values is data, not a
        header — guards the stacked-table detection below."""
        for c in row:
            v = c.value
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                return True
            if isinstance(v, str):
                t = v.strip()
                if any(p.fullmatch(t) for _, p, _ in rules.PROFILE_PATTERNS):
                    return True
        return False

    def _profile_columns(self, ws, skip_cols, header_row, end_row):
        """{column: (entity, header_row_or_0)} for columns whose values
        (between header_row and end_row) overwhelmingly match one
        identifier pattern."""
        found = {}
        start = (header_row or 0) + 1
        if end_row < start:
            return found
        for col_cells in ws.iter_cols(min_row=start, max_row=end_row):
            col = col_cells[0].column
            if col in skip_cols:
                continue
            values = []
            for c in col_cells:
                v = c.value
                if v is None or isinstance(v, bool):
                    continue
                if isinstance(v, (int, float)):
                    v = str(int(v)) if float(v).is_integer() else str(v)
                v = str(v).strip()
                if v and v.lower() not in rules.NULL_MARKERS and not rules.TOKEN_RE.fullmatch(v):
                    values.append(v)
            if len(values) < rules.PROFILE_MIN_SAMPLES:
                continue
            for entity, pat, validator in rules.PROFILE_PATTERNS:
                check = validator or (lambda v: True)
                hits = sum(1 for v in values if pat.fullmatch(v) and check(v))
                if hits / len(values) >= rules.PROFILE_MIN_RATIO:
                    found[col] = (entity, header_row or 0)
                    break
        return found

    def _column_context(self, ws, header_row, max_rows=8, max_words=10):
        """
        Per-column context words for the analyzer, so a bare phone number in
        a column titled 'Phone Number' scores as it would inside a sentence.

        Context must come from labels, never from data: when a header row is
        detected, each column's context is its header cell plus the cell
        directly above it (two-row headers like 'Contact' / 'Mobile'). Only
        when no header row can be found do we fall back to harvesting the
        top rows wholesale.
        """
        words_of = lambda v: [w.lower() for w in re.findall(r"[A-Za-z]{3,}", v)] \
            if isinstance(v, str) else []
        context = {}
        if header_row is not None:
            for cell in ws[header_row]:
                words = words_of(cell.value)
                if header_row > 1:
                    words = words_of(ws.cell(header_row - 1, cell.column).value) + words
                if words:
                    context[cell.column] = words[:max_words]
            return context

        for col_cells in ws.iter_cols(max_row=min(max_rows, ws.max_row)):
            words = []
            for cell in col_cells:
                words.extend(words_of(cell.value))
            if words:
                context[col_cells[0].column] = words[:max_words]
        return context

    def _sheet_segments(self, ws):
        """
        Split a sheet into tables. The first header row is found in the top
        rows; a later row that is header-like, follows an empty row, and
        holds no data-shaped values starts a new table (stacked tables on
        one sheet). Each segment carries its own column rules and context.
        Returns [{header, end, rules, context}] ordered by row.
        """
        first = self._find_header_row(ws)
        headers = []
        if first is not None:
            headers.append(first)
            prev_empty = False
            for row in ws.iter_rows(min_row=first + 1):
                nonempty = [c for c in row if c.value is not None and str(c.value).strip()]
                if not nonempty:
                    prev_empty = True
                    continue
                if (prev_empty and sum(1 for c in row if self._is_headerish(c.value)) >= 2
                        and not self._looks_like_data(row)):
                    headers.append(row[0].row)
                prev_empty = False
        if not headers:
            return [{"header": 0, "end": ws.max_row,
                     "rules": {c: (r[0], "value-profile") for c, r in
                               self._profile_columns(ws, set(), None, ws.max_row).items()},
                     "context": self._column_context(ws, None)}]
        segments = []
        for i, h in enumerate(headers):
            end = headers[i + 1] - 1 if i + 1 < len(headers) else ws.max_row
            seg_rules = {}
            for cell in ws[h]:
                if not self._is_headerish(cell.value):
                    continue
                t = cell.value.strip()
                for entity, pat, deny in rules.COLUMN_RULES:
                    if pat.search(t) and not (deny and deny.search(t)):
                        if entity == "DATE_TIME" and self.dates == "none":
                            break
                        if entity == "LOCATION" and self.locations == "none":
                            break
                        seg_rules[cell.column] = (entity, "column-rule")
                        break
            for col, r in self._profile_columns(ws, set(seg_rules), h, end).items():
                seg_rules[col] = (r[0], "value-profile")
            segments.append({"header": h, "end": end, "rules": seg_rules,
                             "context": self._column_context(ws, h)})
        return segments

    # -------------------------------------------------------------- mask
    def pseudonymize_excel(self, input_path, output_path):
        """Scan every sheet and cell, replacing detected PII spans with tokens."""
        wb = openpyxl.load_workbook(input_path)
        cells_changed = 0

        numeric_entities = rules.NUMERIC_CELL_ENTITIES
        if self.entities is not None:
            numeric_entities = [e for e in numeric_entities if e in self.entities]

        for ws in wb.worksheets:
            segments = self._sheet_segments(ws)
            if not self.column_rules:
                for sg in segments:
                    sg["rules"] = {}
            cache = {}  # (text, numeric?, column, segment) -> (new_text, changed, findings)
            seg_iter = iter(segments)
            sg = next(seg_iter, None)
            for row in ws.iter_rows():
                for cell in row:
                    # advance to the segment containing this row (rows above
                    # the first header belong to no segment)
                    while sg is not None and cell.row > sg["end"]:
                        sg = next(seg_iter, None)
                    in_seg = sg is not None and cell.row > sg["header"]
                    if sg is not None and cell.row == sg["header"]:
                        continue  # header cells are labels, never data
                    value = cell.value
                    entities = None
                    numeric = isinstance(value, (int, float)) and not isinstance(value, bool)

                    # Column rule: mask the whole column below its header —
                    # any value type, no detection needed.
                    rule = sg["rules"].get(cell.column) if in_seg else None
                    if rule and value is not None:
                        rule_text = (str(int(value)) if numeric and float(value).is_integer()
                                     else str(value)).strip()
                        if (rule_text.lower() in rules.NULL_MARKERS
                                or rules.TOKEN_RE.fullmatch(rule_text)):
                            continue
                        token = self.generate_token(rule_text, rule[0])
                        if numeric:
                            self.numeric_tokens.add(token)
                            self.numeric_cells.add(f"{ws.title}!{cell.coordinate}")
                        cell.value = token
                        cells_changed += 1
                        self.report.append(
                            {"where": f"{ws.title}!{cell.coordinate}",
                             "method": rule[1], "entity": rule[0],
                             "score": None, "text": rule_text})
                        continue

                    if isinstance(value, str):
                        text = value
                        # Pre-filter: nothing shorter than 3 chars or without
                        # any letter/digit can be PII — skip the analyzer.
                        stripped = text.strip()
                        if len(stripped) < 3 or not any(c.isalnum() for c in stripped):
                            continue
                    elif numeric:
                        # Card/account numbers stored as numeric cells must
                        # still be analyzed, but only against identifier
                        # patterns — ordinary figures are left untouched.
                        text = str(int(value)) if float(value).is_integer() else str(value)
                        entities = numeric_entities
                        # Pre-filter: every numeric identifier type needs at
                        # least 8 digits (bank 8+, SSN 9, NIC 12, cards 13+),
                        # so shorter numbers (amounts, counts, years) skip
                        # the analyzer entirely.
                        if sum(c.isdigit() for c in text) < rules.MIN_NUMERIC_DIGITS:
                            continue
                    else:
                        continue

                    # Per-run cache: spreadsheets repeat values (cities,
                    # departments, statuses); each distinct (value, entity
                    # set, column context) is analyzed once. Results are
                    # identical by construction — same input, same output.
                    context = sg["context"].get(cell.column) if in_seg else None
                    key = (text, entities is not None,
                           (cell.column, sg["header"]) if context else None)
                    hit = cache.get(key)
                    if hit is None:
                        new_text, changed = self.pseudonymize_text(
                            text, entities=entities, context=context
                        )
                        cache[key] = (new_text, changed, list(self._last_findings))
                    else:
                        new_text, changed, self._last_findings = hit
                    if changed:
                        if entities is not None and rules.TOKEN_RE.fullmatch(new_text):
                            self.numeric_tokens.add(new_text)
                            self.numeric_cells.add(f"{ws.title}!{cell.coordinate}")
                        cell.value = new_text
                        cells_changed += 1
                        for f in self._last_findings:
                            self.report.append(
                                {"where": f"{ws.title}!{cell.coordinate}",
                                 "method": "detected", **f})

        wb.save(output_path)
        print(f"[Success] Excel saved to: {output_path} ({cells_changed} cells modified)")

    # ----------------------------------------------------------- restore
    def depseudonymize_excel(self, input_path, output_path):
        """
        Reverse a masked workbook using the loaded vault. Cells recorded as
        numeric at masking time are restored to numbers; everything else
        stays text, so identifiers like NICs keep their original type.
        """
        wb = openpyxl.load_workbook(input_path)
        cells_restored = 0
        unresolved = 0

        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if not isinstance(cell.value, str) or not rules.TOKEN_RE.search(cell.value):
                        continue
                    # Per-cell record wins (the same value can be a number in
                    # one sheet and text in another); per-token is the
                    # fallback when the vault came from a different workbook.
                    if self.numeric_cells:
                        was_numeric = f"{ws.title}!{cell.coordinate}" in self.numeric_cells
                    else:
                        was_numeric = cell.value.strip() in self.numeric_tokens
                    restored = self.depseudonymize_text(cell.value)
                    if rules.TOKEN_RE.search(restored):
                        unresolved += 1  # token missing from the vault
                    if restored != cell.value:
                        if was_numeric:
                            cell.value = float(restored) if "." in restored else int(restored)
                        else:
                            cell.value = restored
                        cells_restored += 1

        wb.save(output_path)
        print(f"[Success] Restored workbook saved to: {output_path} "
              f"({cells_restored} cells restored)")
        if unresolved:
            print(f"[Warning] {unresolved} cells still contain tokens not found "
                  f"in the vault — was the correct vault loaded?")
