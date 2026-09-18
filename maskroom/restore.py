"""Restore original values inside files an LLM produced from masked input.

The model only ever saw tokens, so a spreadsheet, report or document it
generates carries TOK_<TYPE>_<ID> strings. unmask_file() swaps them back
using the engine's (session) vault with the same tolerant matching as
unmask_text(), and reports what it restored and what it could not resolve.
PDFs are not supported: replacing text inside a PDF reliably is a different
problem.
"""
import io
import os
import re
import zipfile

import openpyxl

from . import rules

TEXT_EXTS = (".md", ".txt", ".csv", ".tsv", ".json", ".html", ".htm", ".xml", ".yaml", ".yml")
EXCEL_EXTS = (".xlsx", ".xlsm")
OFFICE_EXTS = (".docx", ".pptx")
SUPPORTED_EXTS = TEXT_EXTS + EXCEL_EXTS + OFFICE_EXTS

# XML parts of Word/PowerPoint packages that hold user-visible text.
_OFFICE_TEXT_PARTS = re.compile(
    r"^(word/(document|header\d*|footer\d*|footnotes|endnotes|comments)\.xml"
    r"|ppt/(slides|notesSlides|comments)/[^/]+\.xml"
    r"|docProps/core\.xml)$"
)
NOTES_SHEET = "Maskroom notes"


def _merge(total, rep):
    total["restored"] += rep["restored"]
    total["fuzzy"].extend(rep["fuzzy"])
    total["unresolved"].extend(rep["unresolved"])


def _unmask_text_file(engine, in_path, out_path, ext, report):
    with open(in_path, encoding="utf-8", errors="replace") as f:
        text = f.read()
    restored, rep = engine.unmask_text(text)
    _merge(report, rep)
    if rep["unresolved"] and ext in (".md", ".html", ".htm"):
        restored += ("\n\n<!-- maskroom: %d token(s) not in this session's vault: %s -->\n"
                     % (len(rep["unresolved"]), ", ".join(sorted(set(rep["unresolved"])))))
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(restored)


def _unmask_excel(engine, in_path, out_path, report):
    engine.depseudonymize_excel(in_path, out_path)
    wb = openpyxl.load_workbook(out_path)
    restored = 0
    unresolved = []
    src = openpyxl.load_workbook(in_path, read_only=True)
    for ws in src.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str):
                    restored += len(rules.LOOSE_TOKEN_RE.findall(v))  # matches exact tokens too
    src.close()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    unresolved.extend(m.group(0) for m in rules.LEFTOVER_TOKEN_RE.finditer(c.value))
    report["restored"] += max(0, restored - len(unresolved))
    report["unresolved"].extend(unresolved)
    if unresolved:
        ws = wb.create_sheet(NOTES_SHEET)
        ws.append(["Tokens that were not in this session's vault and could not be restored"])
        for t in sorted(set(unresolved)):
            ws.append([t])
        wb.save(out_path)
    wb.close()


def _unmask_office(engine, in_path, out_path, report):
    with zipfile.ZipFile(in_path) as zin, zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if _OFFICE_TEXT_PARTS.match(item.filename):
                xml = data.decode("utf-8", errors="replace")
                # Vault values are plain text; use an escaped view so restored
                # values stay valid XML.
                vault = engine.vault
                engine.vault = vault.escaped()
                try:
                    restored, rep = engine.unmask_text(xml)
                finally:
                    engine.vault = vault
                _merge(report, rep)
                data = restored.encode("utf-8")
            zout.writestr(item, data)


def unmask_file(engine, in_path, out_path):
    """Restore tokens inside a file; returns {"kind", "restored", "fuzzy", "unresolved"}."""
    ext = os.path.splitext(in_path)[1].lower()
    report = {"restored": 0, "fuzzy": [], "unresolved": []}
    if ext in TEXT_EXTS:
        report["kind"] = "text"
        _unmask_text_file(engine, in_path, out_path, ext, report)
    elif ext in EXCEL_EXTS:
        report["kind"] = "excel"
        _unmask_excel(engine, in_path, out_path, report)
    elif ext in OFFICE_EXTS:
        report["kind"] = "office"
        _unmask_office(engine, in_path, out_path, report)
    else:
        raise ValueError(f"unsupported file type {ext!r}; supported: {', '.join(SUPPORTED_EXTS)}")
    report["unresolved"] = sorted(set(report["unresolved"]))
    return report
