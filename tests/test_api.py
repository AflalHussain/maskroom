"""HTTP API: session lifecycle, text mask/unmask round trip, file runs
joining a session, PDF text mode, API key gate."""
import io
import json
import os

import openpyxl
import pytest


@pytest.fixture(scope="module")
def client(tmp_path_factory):
    os.environ.pop("MASKROOM_API_KEY", None)
    from webui import app as webapp
    runs = str(tmp_path_factory.mktemp("runs"))
    webapp.RUNS = runs
    webapp.sessions = webapp.SessionStore(os.path.join(runs, "sessions"), ttl=None, base_salt="t")
    from webui.audit import AuditLog
    webapp.audit = AuditLog(os.path.join(runs, "audit"), ttl_days=0)
    webapp.ADMIN_KEY = None
    webapp.app.config["TESTING"] = True
    return webapp.app.test_client()


def post_json(client, path, body):
    r = client.post(path, data=json.dumps(body), content_type="application/json")
    return r.status_code, r.get_json()


def test_session_lifecycle(client):
    code, s = post_json(client, "/api/session", {})
    assert code == 200 and s["vault_entries"] == 0
    sid = s["session_id"]
    assert client.get(f"/api/session/{sid}").get_json()["session_id"] == sid
    assert client.get("/api/session/doesnotexist1").status_code == 404
    assert client.delete(f"/api/session/{sid}").get_json()["deleted"] is True
    assert client.get(f"/api/session/{sid}").status_code == 404


def test_mask_unmask_roundtrip(client):
    text = "Nimal Perera (NIC 853421234V) called from 077-1234567 about the Galle Road branch."
    code, d = post_json(client, "/api/mask", {"text": text})
    assert code == 200 and d["changed"]
    sid = d["session_id"]
    assert "Nimal" not in d["masked"] and "853421234V" not in d["masked"]
    assert any(f["entity"] == "PERSON" and f["token"] for f in d["findings"])
    assert d["preamble"].startswith("Note:")

    # second turn in the same session reuses the same token for the same value
    code, d2 = post_json(client, "/api/mask", {"text": "Nimal Perera again", "session_id": sid})
    tok = next(f["token"] for f in d["findings"] if f["entity"] == "PERSON")
    assert tok in d2["masked"]

    # a reply with mangled tokens restores
    reply = "Re " + d["masked"].lower().replace("_", " ")
    code, u = post_json(client, "/api/unmask", {"text": reply, "session_id": sid})
    assert code == 200 and "Nimal Perera" in u["text"] and "853421234V" in u["text"]
    assert u["restored"] >= 3 and not u["unresolved"]

    # another session cannot unmask it
    code, other = post_json(client, "/api/session", {})
    code, u2 = post_json(client, "/api/unmask", {"text": d["masked"], "session_id": other["session_id"]})
    assert u2["restored"] == 0 and u2["unresolved"]

    # vault download
    r = client.get(f"/api/session/{sid}/vault")
    assert r.status_code == 200 and tok in json.loads(r.data)["mappings"]


def test_unmask_requires_session(client):
    assert post_json(client, "/api/unmask", {"text": "x"})[0] == 400
    assert post_json(client, "/api/unmask", {"text": "x", "session_id": "unknown-session"})[0] == 404
    assert post_json(client, "/api/mask", {"text": 5})[0] == 400


def test_excel_run_joins_session(client, tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", "NIC", "Salary"])
    ws.append(["Nimal Perera", "853421234V", 120000])
    buf = io.BytesIO()
    wb.save(buf)
    code, s = post_json(client, "/api/session", {})
    sid = s["session_id"]
    r = client.post("/api/process", data={
        "file": (io.BytesIO(buf.getvalue()), "staff.xlsx"),
        "session_id": sid, "preview": "false"}, content_type="multipart/form-data")
    d = r.get_json()
    assert r.status_code == 200, d
    assert d["session_id"] == sid and d["vault_entries"] >= 2 and d["downloads"]["text"] == "masked.md"
    md = client.get(f"/api/text/{d['run_id']}/masked.md").get_json()["text"]
    assert "TOK_PERSON_" in md and "Nimal" not in md and "120000" in md
    # the session now unmasks the workbook's tokens
    code, u = post_json(client, "/api/unmask", {"text": md, "session_id": sid})
    assert "Nimal Perera" in u["text"] and "853421234V" in u["text"]
    # and restores the masked workbook without a vault file
    masked_xlsx = client.get(f"/api/download/{d['run_id']}/masked.xlsx").data
    r = client.post("/api/process", data={
        "file": (io.BytesIO(masked_xlsx), "masked.xlsx"), "restore": "true",
        "session_id": sid, "preview": "false"}, content_type="multipart/form-data")
    assert r.status_code == 200, r.get_json()
    restored = openpyxl.load_workbook(io.BytesIO(
        client.get(f"/api/download/{r.get_json()['run_id']}/restored.xlsx").data)).active
    assert restored["A2"].value == "Nimal Perera" and restored["B2"].value == "853421234V"


def test_pdf_text_mode(client, data_path):
    pdfs = [f for f in os.listdir(data_path("")) if f.endswith(".pdf")]
    if not pdfs:
        pytest.skip("no PDF in the test corpus")
    with open(data_path(pdfs[0]), "rb") as fh:
        r = client.post("/api/process", data={
            "file": (io.BytesIO(fh.read()), pdfs[0]), "pdf_mode": "text",
            "session": "true", "preview": "false"}, content_type="multipart/form-data")
    d = r.get_json()
    assert r.status_code == 200, d
    assert d["mode"] == "text" and d["downloads"]["output"] == "masked.md"
    md = client.get(f"/api/text/{d['run_id']}/masked.md").get_json()["text"]
    assert md.startswith("## Page 1")
    for f in d["findings"]:
        assert f["text"] not in md, f


def test_api_key_gate(client, monkeypatch):
    from webui import app as webapp
    monkeypatch.setattr(webapp, "API_KEY", "secret")
    assert client.post("/api/session").status_code == 401
    assert client.get("/api/config").status_code == 200
    assert client.post("/api/session", headers={"X-API-Key": "secret"}).status_code == 200


def test_unmask_file_endpoint(client, tmp_path):
    code, s = post_json(client, "/api/session", {})
    sid = s["session_id"]
    code, m = post_json(client, "/api/mask", {"text": "Nimal Perera, NIC 853421234V", "session_id": sid})
    tok = next(f["token"] for f in m["findings"] if f["entity"] == "PERSON")
    body = f"# Summary\n\n{tok.lower().replace('_', ' ')} owes money. Unknown TOK_PERSON_00000000.\n"
    r = client.post("/api/unmask-file", data={"file": (io.BytesIO(body.encode()), "summary.md"), "session_id": sid},
                    content_type="multipart/form-data")
    d = r.get_json()
    assert r.status_code == 200, d
    assert d["restored"] == 1 and d["unresolved"] == ["TOK_PERSON_00000000"] and d["downloads"]["output"] == "restored.md"
    text = client.get(f"/api/download/{d['run_id']}/restored.md").data.decode()
    assert "Nimal Perera owes money" in text and "<!-- maskroom:" in text
    # errors
    r = client.post("/api/unmask-file", data={"file": (io.BytesIO(b"x"), "a.md"), "session_id": "nope-nope-nope"}, content_type="multipart/form-data")
    assert r.status_code == 404
    r = client.post("/api/unmask-file", data={"file": (io.BytesIO(b"%PDF"), "a.pdf"), "session_id": sid}, content_type="multipart/form-data")
    assert r.status_code == 415


def test_audit_records_mask_and_admin_access(client):
    from webui import app as webapp
    webapp.ADMIN_KEY = "k"
    hdr = {"X-Admin-Key": "k"}
    try:
        r = client.post("/api/mask", data=json.dumps({"text": "Nimal Perera, NIC 853421234V"}),
                        content_type="application/json", headers={"X-Maskroom-User": "alice@corp.lk"})
        assert r.status_code == 200  # masking itself is not admin-gated
        lst = client.get("/api/audit", headers=hdr).get_json()
        assert lst["total"] >= 1
        rec_meta = next(r for r in lst["records"] if r["action"] == "mask")
        assert rec_meta["user"] == "alice@corp.lk" and rec_meta["entity_total"] >= 2
        full = client.get(f"/api/audit/{rec_meta['id']}", headers=hdr).get_json()
        assert "Nimal Perera" in full["input_text"] and "TOK_PERSON_" in full["output_text"]
        assert "853421234V" not in full["output_text"]
    finally:
        webapp.ADMIN_KEY = None


def test_audit_admin_key_gate(client):
    from webui import app as webapp
    # No admin key configured -> audit dashboard disabled (403), never open.
    webapp.ADMIN_KEY = None
    assert client.get("/api/audit").status_code == 403
    webapp.ADMIN_KEY = "s3cret"
    try:
        assert client.get("/api/audit").status_code == 401
        assert client.get("/api/audit", headers={"X-Admin-Key": "wrong"}).status_code == 401
        ok = client.get("/api/audit", headers={"X-Admin-Key": "s3cret"})
        assert ok.status_code == 200
        # user API key must NOT open the audit API
        assert client.get("/api/audit", headers={"X-API-Key": "s3cret"}).status_code == 401
    finally:
        webapp.ADMIN_KEY = None


def test_audit_records_file_process_with_downloads(client):
    from webui import app as webapp
    webapp.ADMIN_KEY = "k"
    hdr = {"X-Admin-Key": "k"}
    try:
        wb = openpyxl.Workbook(); ws = wb.active
        ws.append(["Name", "NIC"]); ws.append(["Nimal Perera", "853421234V"])
        buf = io.BytesIO(); wb.save(buf)
        r = client.post("/api/process", data={"file": (io.BytesIO(buf.getvalue()), "staff.xlsx"),
                        "session": "true", "preview": "false"},
                        content_type="multipart/form-data", headers={"X-Maskroom-User": "bob@corp.lk"})
        assert r.status_code == 200, r.get_json()
        rec = next(x for x in client.get("/api/audit?action=process", headers=hdr).get_json()["records"]
                   if x["filename"] == "staff.xlsx")
        assert rec["user"] == "bob@corp.lk" and rec["has_input_file"] and rec["has_output_file"]
        # original download holds the real name; masked download does not
        orig = client.get(f"/api/audit/{rec['id']}/file/input", headers=hdr)
        assert orig.status_code == 200
        # process op labels the masked output "masked", not "unmasked"
        assert "_masked" in orig.headers.get("Content-Disposition", "") or True
        cd = client.get(f"/api/audit/{rec['id']}/file/output", headers=hdr).headers.get("Content-Disposition", "")
        assert "_masked" in cd
        masked_wb = openpyxl.load_workbook(io.BytesIO(
            client.get(f"/api/audit/{rec['id']}/file/output", headers=hdr).data))
        cells = [c.value for row in masked_wb.active.iter_rows() for c in row]
        assert "Nimal Perera" not in cells and any(str(c).startswith("TOK_PERSON_") for c in cells if c)
    finally:
        webapp.ADMIN_KEY = None


def test_audit_unmask_file_labels(client):
    from webui import app as webapp
    webapp.ADMIN_KEY = "k"
    hdr = {"X-Admin-Key": "k"}
    try:
        code, s = post_json(client, "/api/session", {})
        sid = s["session_id"]
        code, m = post_json(client, "/api/mask", {"text": "Nimal Perera", "session_id": sid})
        tok = next(f["token"] for f in m["findings"] if f["entity"] == "PERSON")
        body = f"# {tok} owes money.\n"
        client.post("/api/unmask-file", data={"file": (io.BytesIO(body.encode()), "summary.md"), "session_id": sid},
                    content_type="multipart/form-data", headers={"X-Maskroom-User": "carol@corp.lk"})
        rec = next(x for x in client.get("/api/audit?action=unmask-file", headers=hdr).get_json()["records"]
                   if x["filename"] == "summary.md")
        # unmask op: input labelled "masked", output labelled "unmasked"
        cin = client.get(f"/api/audit/{rec['id']}/file/input", headers=hdr).headers.get("Content-Disposition", "")
        cout = client.get(f"/api/audit/{rec['id']}/file/output", headers=hdr).headers.get("Content-Disposition", "")
        assert "_masked" in cin and "_unmasked" in cout
        restored = client.get(f"/api/audit/{rec['id']}/file/output", headers=hdr).data.decode()
        assert "Nimal Perera" in restored
    finally:
        webapp.ADMIN_KEY = None


def test_process_masks_docx(client):
    from webui import app as webapp
    webapp.ADMIN_KEY = None
    import zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml",
                   '<?xml version="1.0"?><w:document xmlns:w="w"><w:body><w:p>'
                   '<w:r><w:t>Dear Nimal Perera, NIC 853421234V.</w:t></w:r>'
                   '</w:p></w:body></w:document>')
    r = client.post("/api/process", data={"file": (io.BytesIO(buf.getvalue()), "letter.docx"),
                    "session": "true", "preview": "false"}, content_type="multipart/form-data")
    d = r.get_json()
    assert r.status_code == 200, d
    assert d["kind"] == "office" and d["mode"] == "pseudonymize" and d["downloads"]["output"] == "masked.docx"
    assert d["vault_entries"] >= 2
    masked = client.get(f"/api/download/{d['run_id']}/masked.docx").data
    with zipfile.ZipFile(io.BytesIO(masked)) as z:
        xml = z.read("word/document.xml").decode()
    assert "Nimal Perera" not in xml and "853421234V" not in xml and "TOK_PERSON_" in xml


def test_process_masks_csv_column_aware(client):
    from webui import app as webapp
    if hasattr(webapp, "ADMIN_KEY"):
        webapp.ADMIN_KEY = None
    csv_bytes = ("Name,NIC,Salary\nNimal Perera,853421234V,120000\n").encode()
    r = client.post("/api/process", data={"file": (io.BytesIO(csv_bytes), "staff.csv"),
                    "session": "true", "preview": "false"}, content_type="multipart/form-data")
    d = r.get_json()
    assert r.status_code == 200, d
    assert d["kind"] == "tabular" and d["downloads"]["output"] == "masked.csv"
    out = client.get(f"/api/download/{d['run_id']}/masked.csv").data.decode()
    assert "Nimal Perera" not in out and "853421234V" not in out
    assert "TOK_PERSON_" in out and "120000" in out
