"""Per-column MASK/KEEP scoring of the hostile stress workbook against its
answer key, plus free-text leak/keep checks. One known ceiling is tolerated
for the default (lg) model: the Sinhala-headed name column depends on NER."""
import collections
import json
import re

import openpyxl
import pytest
from openpyxl.utils import column_index_from_string as cidx

from maskroom.rules import TOKEN_RE

KNOWN_LG_CEILING = {"Sinhala Headers|B"}  # header 'නම' -> NER only (lg ~92%)


def _score(key, original, masked, report):
    methods = collections.defaultdict(collections.Counter)
    for e in report:
        sh, coord = e["where"].split("!")
        col = re.match(r"[A-Z]+", coord).group(0)
        methods[f"{sh}|{col}"][e["method"]] += 1
    fails = []
    for k, exp in key["columns"].items():
        sh, col = k.split("|")
        c = cidx(col)
        wo, wm = original[sh], masked[sh]
        rows = range(17 if sh == "Deep Header" else 2, wo.max_row + 1)
        if exp == "MIXED":
            if col == "A":
                segs = [("MASK", range(2, 8)), ("KEEP", range(11, wo.max_row + 1))]
            else:
                segs = [("KEEP", range(2, 8)), ("MASK", range(11, wo.max_row + 1))]
        else:
            segs = [(exp, rows)]
        for seg_exp, seg_rows in segs:
            tot = hit = 0
            for r in seg_rows:
                if sh == "Stacked Tables" and r == 10:
                    continue  # second table's header row
                v = wo.cell(r, c).value
                if v is None or str(v).strip().lower() in ("", "n/a", "-"):
                    continue
                tot += 1
                mv = wm.cell(r, c).value
                if isinstance(mv, str) and TOKEN_RE.search(mv):
                    hit += 1
            ok = (hit == tot) if seg_exp == "MASK" else (hit == 0)
            if not ok:
                fails.append((k, seg_exp, hit, tot, dict(methods.get(k, {}))))
    return fails


@pytest.fixture(scope="module")
def stress_run(make_engine, tmp_path_factory):
    import os
    data = os.path.join(os.path.dirname(__file__), "data")
    out = tmp_path_factory.mktemp("stress") / "masked.xlsx"
    e = make_engine()
    e.pseudonymize_excel(os.path.join(data, "PII_Stress_Test.xlsx"), str(out))
    key = json.load(open(os.path.join(data, "PII_Stress_Test_key.json")))
    return (key, openpyxl.load_workbook(os.path.join(data, "PII_Stress_Test.xlsx")),
            openpyxl.load_workbook(out), e.report)


def test_columns_match_answer_key(stress_run, nlp_model):
    key, original, masked, report = stress_run
    fails = _score(key, original, masked, report)
    tolerated = set() if nlp_model else KNOWN_LG_CEILING
    hard = [f for f in fails if f[0] not in tolerated]
    assert not hard, f"column checks failed: {hard}"


def test_free_text_leaks_and_keeps(stress_run):
    key, _, masked, _ = stress_run
    notes = " ".join(str(masked["Free Text"].cell(r, 2).value) for r in range(2, 8))
    assert [x for x in key["free_text_must_mask"] if x in notes] == []
    assert [x for x in key["free_text_must_keep"] if x not in notes] == []


def test_methods_present(stress_run):
    _, _, _, report = stress_run
    counts = collections.Counter(e["method"] for e in report)
    assert counts["column-rule"] > 0 and counts["value-profile"] > 0 and counts["detected"] > 0
