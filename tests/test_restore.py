"""Restoring tokens inside files an LLM produced (maskroom/restore.py)."""
import os
import zipfile

import openpyxl
import pytest

from maskroom.restore import NOTES_SHEET, unmask_file


@pytest.fixture
def vault_engine(make_engine):
    e = make_engine()
    e.vault.update_from_dict({"TOK_PERSON_8B584CCF": "Nimal Perera", "TOK_LK_NIC_0F0F0F0F": "853421234V",
                              "TOK_EMAIL_ADDRESS_ABCDEF12": "a&b@x.lk"})
    return e


def test_markdown_roundtrip_with_marker(vault_engine, tmp_path):
    src = tmp_path / "report.md"
    src.write_text("# Report\n\n| who | nic |\n|---|---|\n| tok person 8b584ccf | TOK\\_LK\\_NIC\\_0F0F0F0F |\n\nUnknown TOK_PERSON_99999999.\n")
    out = tmp_path / "restored.md"
    rep = unmask_file(vault_engine, str(src), str(out))
    text = out.read_text()
    assert rep["kind"] == "text" and rep["restored"] == 2 and len(rep["fuzzy"]) == 2
    assert "Nimal Perera" in text and "853421234V" in text
    assert rep["unresolved"] == ["TOK_PERSON_99999999"] and "<!-- maskroom: 1 token(s)" in text


def test_csv_has_no_marker(vault_engine, tmp_path):
    src = tmp_path / "x.csv"; src.write_text("name,id\nTOK_PERSON_8B584CCF,TOK_PERSON_99999999\n")
    out = tmp_path / "r.csv"
    rep = unmask_file(vault_engine, str(src), str(out))
    assert out.read_text() == "name,id\nNimal Perera,TOK_PERSON_99999999\n" and rep["unresolved"]


def test_excel_roundtrip_and_notes_sheet(vault_engine, tmp_path):
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Customer", "NIC", "Amount"])
    ws.append(["TOK_PERSON_8B584CCF", "tok lk nic 0f0f0f0f", 120000])
    ws.append(["TOK_PERSON_DEADBEEF", "", 5])
    src = tmp_path / "out.xlsx"; wb.save(src)
    out = tmp_path / "restored.xlsx"
    rep = unmask_file(vault_engine, str(src), str(out))
    r = openpyxl.load_workbook(out)
    assert r.active["A2"].value == "Nimal Perera" and r.active["B2"].value == "853421234V"
    assert r.active["C2"].value == 120000
    assert rep["kind"] == "excel" and rep["restored"] == 2 and rep["unresolved"] == ["TOK_PERSON_DEADBEEF"]
    assert NOTES_SHEET in r.sheetnames and r[NOTES_SHEET]["A2"].value == "TOK_PERSON_DEADBEEF"


def test_excel_without_unresolved_has_no_notes_sheet(vault_engine, tmp_path):
    wb = openpyxl.Workbook(); wb.active.append(["TOK_PERSON_8B584CCF"])
    src = tmp_path / "o.xlsx"; wb.save(src); out = tmp_path / "r.xlsx"
    unmask_file(vault_engine, str(src), str(out))
    assert NOTES_SHEET not in openpyxl.load_workbook(out).sheetnames


def test_docx_text_parts_restored_and_escaped(vault_engine, tmp_path):
    src = tmp_path / "letter.docx"
    doc_xml = ('<?xml version="1.0"?><w:document xmlns:w="w"><w:body><w:p><w:r><w:t>Dear TOK_PERSON_8B584CCF,'
               '</w:t></w:r><w:r><w:t>mail TOK_EMAIL_ADDRESS_ABCDEF12</w:t></w:r></w:p></w:body></w:document>')
    with zipfile.ZipFile(src, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", doc_xml)
        z.writestr("word/media/image1.png", b"\x89PNGbinary")
    out = tmp_path / "restored.docx"
    rep = unmask_file(vault_engine, str(src), str(out))
    with zipfile.ZipFile(out) as z:
        xml = z.read("word/document.xml").decode()
        assert z.read("word/media/image1.png") == b"\x89PNGbinary"
    assert "Dear Nimal Perera," in xml and "a&amp;b@x.lk" in xml
    assert rep["kind"] == "office" and rep["restored"] == 2 and not rep["unresolved"]


def test_unsupported_type(vault_engine, tmp_path):
    src = tmp_path / "x.pdf"; src.write_bytes(b"%PDF-1.4")
    with pytest.raises(ValueError):
        unmask_file(vault_engine, str(src), str(tmp_path / "r.pdf"))
