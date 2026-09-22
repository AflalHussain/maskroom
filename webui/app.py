"""Web UI + text/JSON API for the PII masking engine.
Run:  pii_env/bin/python webui/app.py

Pages:   /            file masking studio (Excel/PDF, preview, vault download)
         /staging     LLM staging area: mask text and files, copy them into
                      Claude (or any chat app), paste the reply back to unmask
API:     see docs in README.md "LLM staging" — /api/session, /api/mask,
         /api/unmask, /api/process, /api/download

Engines are built once per option set and share one spaCy load; all analysis
is serialized through the session store's engine lock (spaCy pipelines are
not guaranteed thread-safe). Authentication lives in webui/auth.py: single
sign-on (MASKROOM_AUTH_MODE=oidc) with roles and service keys, or the legacy
shared secrets with MASKROOM_AUTH_MODE=off.
"""
import base64
import io
import json
import os
import sys
import threading
import time
import uuid

import openpyxl
from flask import Flask, Response, jsonify, request, send_file, send_from_directory

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from maskroom import FinancialPrivacyEngine, SessionStore, build_nlp_engine
from maskroom.store import PolicyStore, connect as db_connect
from maskroom import overlay as overlay_mod
from maskroom import rules
from maskroom.locale import DEFAULT_LOCALE, available_locales
from maskroom.pipeline import SUPPORTED_EXTS as MASK_EXTS, mask_file
from maskroom.restore import SUPPORTED_EXTS as RESTORE_EXTS, unmask_file
from webui.audit import AuditLog
from webui import auth as auth_mod
from webui.auth import (configure_auth, current_principal, external_base as _external_base,
                        client_ip as _ip, owned)

try:
    import pymupdf as fitz
except ImportError:
    import fitz

BASE = os.path.dirname(os.path.abspath(__file__))
REPO_ROOT = os.path.dirname(BASE)
# Files live under MASKROOM_DATA_DIR (runs/, audit/, ext/, the default SQLite
# db); without it the legacy layout inside the source tree is used.
DATA_DIR = os.environ.get("MASKROOM_DATA_DIR")
RUNS = os.path.join(DATA_DIR, "runs") if DATA_DIR else os.path.join(BASE, "runs")
AUDIT_DIR = os.path.join(DATA_DIR, "audit") if DATA_DIR else os.path.join(RUNS, "audit")
os.makedirs(RUNS, exist_ok=True)

MAX_PREVIEW_ROWS = 80
MAX_PREVIEW_COLS = 14
MAX_PREVIEW_PAGES = 8
MAX_TEXT_CHARS = int(os.environ.get("MAX_TEXT_CHARS", "200000"))
MAX_MARKDOWN_ROWS = 2000
# Legacy shared secrets (auth mode "off" only); read live by webui.auth so
# tests can swap them on this module.
API_KEY = os.environ.get("MASKROOM_API_KEY")
ADMIN_KEY = os.environ.get("MASKROOM_ADMIN_KEY")
TTL_HOURS = float(os.environ.get("SESSION_TTL_HOURS", "24"))
AUDIT_TTL_DAYS = float(os.environ.get("AUDIT_TTL_DAYS", "90"))  # 0 = keep forever
EXT_DIR = os.path.join(REPO_ROOT, "extension")
# Directory holding the packaged maskroom-<version>.crx for self-hosted install.
EXT_DIST_DIR = os.environ.get(
    "EXT_DIST_DIR", os.path.join(DATA_DIR, "ext") if DATA_DIR else os.path.join(REPO_ROOT, "local"))

app = Flask(__name__, static_folder="static")
# Runtime state (sessions, audit trail, admin rules) lives in the database:
# MASKROOM_DATABASE_URL, or SQLite under the data dir. Opened at import so a
# missing database fails the process fast rather than the first request.
_db = db_connect()
sessions = SessionStore(_db, ttl=TTL_HOURS * 3600 or None)
audit = AuditLog(_db, AUDIT_DIR, ttl_days=AUDIT_TTL_DAYS)
_policy = PolicyStore(_db)

# Admin pages are always served; their JS asks /api/me and shows the sign-in
# or "role required" state. The APIs behind them are gated in webui/auth.py.
ADMIN_PAGES = {"/admin": "admin.html", "/admin/rules": "rules.html", "/admin/users": "users.html"}


def _user():
    """Who to attribute this request to in the audit trail: the signed-in
    user's email or the service key's name. With auth off, the header the
    extension used to send is still honoured."""
    p = current_principal()
    if p.authenticated:
        return p.label
    return request.headers.get("X-Maskroom-User") or "unknown"


def _entity_counts(findings):
    by = {}
    for f in findings or []:
        by[f["entity"]] = by.get(f["entity"], 0) + 1
    return by


def _audit(**kw):
    """Record an audit entry; never let auditing break the request."""
    try:
        audit.maybe_sweep()
        audit.record(**kw)
    except Exception as e:  # noqa: BLE001
        app.logger.warning("audit record failed: %s", e)

# ---------------------------------------------------------------- engines
_nlp_engines = {}
_engines = {}
_build_lock = threading.Lock()

# Admin policy overlay (custom deny/allow/regex). The fingerprint is folded
# into the engine cache key so a save via /api/policy hot-swaps every warm
# engine. Other processes (gunicorn workers, a second instance, the admin CLI)
# save to the same database; _refresh_overlay polls the revision id so their
# changes take effect here within _OVERLAY_POLL_S. Guarded by _build_lock.
_overlay, _overlay_rev = _policy.current()
_overlay_fp = overlay_mod.fingerprint(_overlay)
_overlay_checked = time.monotonic()
_OVERLAY_POLL_S = 5.0


def _refresh_overlay(force=False):
    """Reload the overlay if the database has a newer revision than the one
    the warm engines were built from (checked at most every few seconds)."""
    global _overlay, _overlay_fp, _overlay_rev, _overlay_checked
    now = time.monotonic()
    if not force and now - _overlay_checked < _OVERLAY_POLL_S:
        return
    with _build_lock:
        if not force and now - _overlay_checked < _OVERLAY_POLL_S:
            return
        _overlay_checked = now
        if force or _policy.latest_revision_id() != _overlay_rev:
            _overlay, _overlay_rev = _policy.current()
            _overlay_fp = overlay_mod.fingerprint(_overlay)
            _engines.clear()


def _bool(v, default=True):
    if v is None or v == "":
        return default
    return str(v).lower() not in ("false", "0", "no", "off")


def engine_options(opts):
    """Normalized engine options from a form or JSON mapping."""
    entities = opts.get("entities") or ""
    if isinstance(entities, str):
        entities = [e.strip() for e in entities.split(",") if e.strip()]
    return {
        "min_score": float(opts.get("min_score") or 0.6),
        "entities": tuple(entities) or None,
        "nlp_model": opts.get("nlp_model") or None,
        "dates": opts.get("dates") or "birth",
        "locations": opts.get("locations") or "address",
        "locale": opts.get("locale") or DEFAULT_LOCALE,
        "column_rules": _bool(opts.get("column_rules"), True),
    }


def get_engine(options):
    """A warm engine for this option set. Engines share the spaCy model
    for their nlp_model, so a new option combination costs only recognizer
    setup, not a model load."""
    _refresh_overlay()
    overlay, fp = _overlay, _overlay_fp  # one snapshot, so a refresh can't mix them
    key = tuple(sorted(options.items())) + (("overlay", fp),)
    eng = _engines.get(key)
    if eng is not None:
        return eng
    with _build_lock:
        eng = _engines.get(key)
        if eng is None:
            model = options["nlp_model"]
            nlp = _nlp_engines.get(model)
            if nlp is None:
                nlp = _nlp_engines[model] = build_nlp_engine(model)
            kw = dict(options, entities=list(options["entities"]) if options["entities"] else None)
            eng = _engines[key] = FinancialPrivacyEngine(nlp_engine=nlp, overlay=overlay, **kw)
    return eng


# ------------------------------------------------ self-hosted extension
import glob as _glob
import re as _re
import zipfile as _zipfile


def _latest_crx():
    """Path to the newest maskroom .crx in EXT_DIST_DIR, or None."""
    crxs = sorted(_glob.glob(os.path.join(EXT_DIST_DIR, "*.crx")), key=os.path.getmtime)
    # Absolute: send_file resolves relative paths against the Flask app root, not CWD.
    return os.path.abspath(crxs[-1]) if crxs else None


def _crx_version(path):
    """Read the extension version from inside the .crx (CRX3 = 'Cr24' + u32
    version + u32 header length + header + ZIP). Falls back to the filename."""
    try:
        with open(path, "rb") as f:
            magic = f.read(4)
            if magic == b"Cr24":
                import struct
                f.read(4)  # format version
                (hlen,) = struct.unpack("<I", f.read(4))
                f.seek(12 + hlen)
                zbytes = io.BytesIO(f.read())
            else:
                f.seek(0); zbytes = io.BytesIO(f.read())
        with _zipfile.ZipFile(zbytes) as z:
            return json.loads(z.read("manifest.json").decode("utf-8"))["version"]
    except Exception:  # noqa: BLE001
        m = _re.search(r"(\d+(?:\.\d+)+)", os.path.basename(path))
        return m.group(1) if m else "0.0.0"


def _extension_id():
    """Compute the extension ID from the manifest 'key' (base64 DER public
    key): sha256(DER)[:16] mapped 0-f -> a-p. Matches Chrome's algorithm."""
    try:
        key = json.load(open(os.path.join(EXT_DIR, "manifest.json")))["key"]
        der = base64.b64decode(key)
        import hashlib
        h = hashlib.sha256(der).hexdigest()[:32]
        return h.translate(str.maketrans("0123456789abcdef", "abcdefghijklmnop"))
    except Exception:  # noqa: BLE001
        return None


# Authentication and the access matrix: webui/auth.py. Configured here so the
# legacy keys are read from this module (tests set webapp.API_KEY/ADMIN_KEY).
configure_auth(app, _db, legacy_key=lambda: API_KEY, admin_key=lambda: ADMIN_KEY,
               extension_ids=[x for x in (_extension_id(),) if x])
app.before_request(auth_mod.authenticate)


@app.after_request
def no_store(resp):
    if request.path.startswith("/api/"):
        resp.headers["Cache-Control"] = "no-store"
    return resp


# ----------------------------------------------------------------- pages
@app.get("/")
def index():
    return send_from_directory(os.path.join(BASE, "static"), "index.html")


@app.get("/staging")
def staging():
    return send_from_directory(os.path.join(BASE, "static"), "staging.html")


@app.get("/admin")
def admin_page():
    return send_from_directory(os.path.join(BASE, "static"), "admin.html")


@app.get("/admin/rules")
def rules_page():
    return send_from_directory(os.path.join(BASE, "static"), "rules.html")


@app.get("/admin/users")
def admin_users_page():
    return send_from_directory(os.path.join(BASE, "static"), "users.html")


# ---- self-hosted extension install (open; Chrome sends no auth) ----
@app.get("/ext/update.xml")
def ext_update_xml():
    """Update manifest for force-installed self-hosting. The codebase URL is
    built from THIS request, so it is correct on ngrok/prod/localhost without
    editing anything. Point the ExtensionSettings update_url at this route."""
    crx = _latest_crx()
    ext_id = _extension_id()
    if not crx or not ext_id:
        return Response("<!-- no .crx in EXT_DIST_DIR or no manifest key -->",
                        status=503, mimetype="application/xml")
    codebase = f"{_external_base()}/ext/maskroom.crx"
    xml = ('<?xml version="1.0" encoding="UTF-8"?>\n'
           '<gupdate xmlns="http://www.google.com/update2/response" protocol="2.0">\n'
           f'  <app appid="{ext_id}">\n'
           f'    <updatecheck codebase="{codebase}" version="{_crx_version(crx)}"/>\n'
           '  </app>\n</gupdate>\n')
    return Response(xml, mimetype="application/xml",
                    headers={"Cache-Control": "no-store"})


@app.get("/ext/maskroom.crx")
def ext_crx():
    crx = _latest_crx()
    if not crx:
        return jsonify({"error": "No packaged extension available."}), 404
    return send_file(crx, mimetype="application/x-chrome-extension",
                     as_attachment=True,
                     download_name=f"maskroom-{_crx_version(crx)}.crx")


@app.get("/api/config")
def config():
    oidc = auth_mod.AUTH.mode == "oidc"
    return jsonify({"auth_mode": auth_mod.AUTH.mode, "login_url": "/auth/login",
                    "auth_required": oidc or bool(API_KEY), "default_locale": DEFAULT_LOCALE,
                    "locales": available_locales(), "session_ttl_hours": TTL_HOURS,
                    "preamble": rules.LLM_TOKEN_PREAMBLE, "max_text_chars": MAX_TEXT_CHARS,
                    "restore_exts": list(RESTORE_EXTS),
                    "admin_auth_required": oidc or bool(ADMIN_KEY),
                    "audit_ttl_days": AUDIT_TTL_DAYS})


@app.get("/api/locales")
def locales():
    return jsonify({"default": DEFAULT_LOCALE, "locales": available_locales()})


# -------------------------------------------------------------- sessions
def _owner():
    """The principal id new sessions and runs are attributed to."""
    return current_principal().id


def _session_or_error(session_id, create=False):
    """(session, error_response). Sweeps idle sessions as a side effect. A
    session another principal owns is reported as unknown, not forbidden."""
    sessions.maybe_sweep()
    if not session_id:
        if create:
            return sessions.create(owner_id=_owner()), None
        return None, (jsonify({"error": "session_id is required."}), 400)
    sess = sessions.get(session_id)
    if sess is None or not owned(sess.owner_id):
        return None, (jsonify({"error": "Unknown or expired session."}), 404)
    return sess, None


@app.post("/api/session")
def session_create():
    sessions.maybe_sweep()
    return jsonify(sessions.create(owner_id=_owner()).info())


@app.get("/api/session/<session_id>")
def session_info(session_id):
    sess, err = _session_or_error(session_id)
    return err or jsonify(sess.info())


@app.delete("/api/session/<session_id>")
def session_delete(session_id):
    sess = sessions.get(session_id)
    if sess is not None and not owned(sess.owner_id):
        return jsonify({"deleted": False})
    return jsonify({"deleted": sessions.delete(session_id)})


@app.get("/api/session/<session_id>/vault")
def session_vault(session_id):
    sess, err = _session_or_error(session_id)
    if err:
        return err
    with sess.lock:
        doc = sess.export()
    body = json.dumps(doc, ensure_ascii=False, indent=2).encode("utf-8")
    return send_file(io.BytesIO(body), mimetype="application/json", as_attachment=True,
                     download_name=f"vault-{session_id}.json")


# ------------------------------------------------------------------ text
def _json_body():
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return None, (jsonify({"error": "Send a JSON object body."}), 400)
    text = data.get("text")
    if not isinstance(text, str):
        return None, (jsonify({"error": '"text" (string) is required.'}), 400)
    if len(text) > MAX_TEXT_CHARS:
        return None, (jsonify({"error": f"Text is longer than {MAX_TEXT_CHARS} characters."}), 413)
    return data, None


@app.post("/api/mask")
def mask_text():
    """{text, session_id?, min_score?, dates?, locations?, locale?, entities?,
    nlp_model?} -> {session_id, masked, changed, findings, vault_entries,
    preamble, elapsed_s}. Without a session_id a new session is created."""
    data, err = _json_body()
    if err:
        return err
    sess, err = _session_or_error(data.get("session_id"), create=True)
    if err:
        return err
    try:
        engine = get_engine(engine_options(data))
    except Exception as e:
        return jsonify({"error": str(e)}), 400

    t0 = time.time()
    with sessions.bind(engine, sess):
        masked, changed = engine.pseudonymize_text(data["text"])
        findings = list(engine._last_findings)
        # token_for is O(vault) per finding, so this loop is O(vault * findings).
        # Fine at realistic session sizes (hundreds of tokens, a few findings);
        # sub-millisecond. If sessions ever accumulate many thousands of
        # pseudonyms and this shows up in profiling, add a Vault.inverse() that
        # builds the reverse map once and look up against it here (O(vault + findings)).
        for f in findings:
            f["token"] = engine.vault.token_for(f["text"].strip())
        entries = len(engine.vault)
    _audit(action="mask", user=_user(), session_id=sess.id, ip=_ip(), kind="text",
           input_text=data["text"], output_text=masked, by_entity=_entity_counts(findings),
           entity_total=len(findings), changed=changed)
    return jsonify({
        "session_id": sess.id, "masked": masked, "changed": changed,
        "findings": findings, "vault_entries": entries,
        "preamble": rules.LLM_TOKEN_PREAMBLE, "elapsed_s": round(time.time() - t0, 2),
    })


@app.post("/api/unmask")
def unmask_text():
    """{text, session_id} -> {text, restored, fuzzy, unresolved}."""
    data, err = _json_body()
    if err:
        return err
    sess, err = _session_or_error(data.get("session_id"))
    if err:
        return err
    engine = get_engine(engine_options({}))
    with sessions.bind(engine, sess):
        restored, report = engine.unmask_text(data["text"])
    _audit(action="unmask", user=_user(), session_id=sess.id, ip=_ip(), kind="text",
           input_text=data["text"], output_text=restored,
           entity_total=report.get("restored", 0), unresolved=report.get("unresolved", []))
    return jsonify({"session_id": sess.id, "text": restored, **report})


# ----------------------------------------------------------------- files
def excel_preview(path, limit_sheets=6):
    """Sheets as row-lists of strings, capped for the browser."""
    wb = openpyxl.load_workbook(path, read_only=True)
    sheets = []
    for name in wb.sheetnames[:limit_sheets]:
        ws = wb[name]
        rows = []
        for r, row in enumerate(ws.iter_rows(values_only=True)):
            if r >= MAX_PREVIEW_ROWS:
                break
            rows.append(["" if v is None else str(v) for v in row[:MAX_PREVIEW_COLS]])
        sheets.append({"name": name, "rows": rows,
                       "truncated": ws.max_row > MAX_PREVIEW_ROWS
                                    or (ws.max_column or 0) > MAX_PREVIEW_COLS})
    wb.close()
    return sheets


def excel_markdown(path, out_path):
    """Render a (masked) workbook as Markdown tables, one section per sheet,
    so it can be pasted into a chat without needing code execution."""
    wb = openpyxl.load_workbook(path, read_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"## Sheet: {ws.title}\n")
        rows = []
        for r, row in enumerate(ws.iter_rows(values_only=True)):
            if r >= MAX_MARKDOWN_ROWS:
                parts.append(f"\n_(truncated after {MAX_MARKDOWN_ROWS} rows)_\n")
                break
            cells = ["" if v is None else str(v).replace("|", "\\|").replace("\n", " ")
                     for v in row]
            if any(cells):
                rows.append("| " + " | ".join(cells) + " |")
        if rows:
            width = rows[0].count("|") - 1
            parts.append(rows[0])
            parts.append("|" + "---|" * width)
            parts.extend(rows[1:])
        parts.append("")
    wb.close()
    text = "\n".join(parts)
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(text)
    return text


def pdf_preview(path, dpi=90):
    doc = fitz.open(path)
    pages = []
    for i in range(min(len(doc), MAX_PREVIEW_PAGES)):
        png = doc[i].get_pixmap(dpi=dpi).tobytes("png")
        pages.append(base64.b64encode(png).decode())
    n = len(doc)
    doc.close()
    return {"pages": pages, "total": n}


@app.post("/api/process")
def process():
    """Multipart: file (+ options). Optional session_id joins the run to a
    session vault; pdf_mode=text|redact (default redact); restore=true with
    either a vault file or a session_id reverses a masked workbook;
    preview=false skips the before/after previews."""
    f = request.files.get("file")
    if not f or not f.filename:
        return jsonify({"error": "No file uploaded."}), 400
    name = os.path.basename(f.filename)
    ext = os.path.splitext(name)[1].lower()
    if ext not in MASK_EXTS:
        return jsonify({"error": f"Unsupported file type {ext!r} — use {', '.join(MASK_EXTS)}."}), 400
    is_office = ext in (".docx", ".pptx")
    is_csv = ext in (".csv", ".tsv")
    is_textfile = ext in (".txt", ".json")

    opts = request.form
    restore = opts.get("restore") == "true"
    pdf_text = ext == ".pdf" and opts.get("pdf_mode", "redact") == "text"
    want_preview = _bool(opts.get("preview"), True)
    session_id = opts.get("session_id") or None
    sess = None
    if session_id or opts.get("session") == "true":
        sess, err = _session_or_error(session_id, create=not restore)
        if err:
            return err
    try:
        engine = get_engine(engine_options(opts))
    except Exception as e:
        return jsonify({"error": str(e)}), 400

    run_id = uuid.uuid4().hex[:12]
    run_dir = os.path.join(RUNS, run_id)
    os.makedirs(run_dir)
    auth_mod.AUTH.runs.create(run_id, owner_id=_owner(), session_id=sess.id if sess else None)
    in_path = os.path.join(run_dir, "input" + ext)
    f.save(in_path)

    if restore:
        out_name = "restored" + ext
    elif pdf_text:
        out_name = "masked.md"
    else:
        out_name = ("masked" if ext != ".pdf" else "redacted") + ext
    out_path = os.path.join(run_dir, out_name)
    text_name = None
    t0 = time.time()

    # A run always has a scratch session so per-run state (vault/report) is
    # isolated; a real session makes the vault shared across the conversation.
    scratch = sess or sessions.create(owner_id=_owner())
    kind = mode = None
    try:
        with sessions.bind(engine, scratch):
            if restore:
                if ext not in (".xlsx", ".xlsm"):
                    return jsonify({"error": "Restore here only works for Excel; use /api/unmask-file for docx/pptx/text."}), 400
                vf = request.files.get("vault")
                if vf:
                    vpath = os.path.join(run_dir, "vault_in.json")
                    vf.save(vpath)
                    engine.load_vault(vpath)
                elif not sess:
                    return jsonify({"error": "Restore needs the vault JSON saved when the "
                                             "file was masked, or the session it was masked in."}), 400
                engine.depseudonymize_excel(in_path, out_path)
                findings = list(engine.report)
                kind, mode = "excel", "restore"
            else:
                result = mask_file(engine, in_path, out_path, pdf_text=pdf_text)
                findings, kind, mode = result.findings, result.kind, result.mode
                if result.kind == "excel":
                    text_name = "masked.md"
                    excel_markdown(out_path, os.path.join(run_dir, text_name))
                engine.save_vault(os.path.join(run_dir, "vault.json"))
            vault_entries = len(engine.vault)
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    except Exception as e:
        return jsonify({"error": f"Processing failed: {e}"}), 500
    finally:
        if scratch is not sess:
            sessions.delete(scratch.id)
    elapsed = time.time() - t0

    by_entity = {}
    for e in findings:
        by_entity[e["entity"]] = by_entity.get(e["entity"], 0) + 1

    resp = {
        "run_id": run_id,
        "session_id": sess.id if sess else None,
        "filename": name,
        "kind": kind,
        "mode": mode,
        "elapsed_s": round(elapsed, 1),
        "vault_entries": vault_entries,
        "findings": findings,
        "by_entity": dict(sorted(by_entity.items(), key=lambda kv: -by_entity[kv[0]])),
        "downloads": {"output": out_name,
                      "vault": "vault.json" if not restore else None,
                      "text": out_name if pdf_text else text_name},
    }
    if want_preview and not is_office and not is_csv and not is_textfile:
        if ext == ".pdf":
            resp["before"] = pdf_preview(in_path)
            resp["after"] = pdf_preview(out_path) if not pdf_text else None
        else:
            resp["before"] = excel_preview(in_path)
            resp["after"] = excel_preview(out_path)

    with open(os.path.join(run_dir, "result.json"), "w") as fh:
        json.dump({k: v for k, v in resp.items() if k not in ("before", "after")}, fh)
    if not restore:
        _audit(action="process", user=_user(), session_id=sess.id if sess else None, ip=_ip(),
               kind=resp["kind"], filename=name, input_file=in_path, output_file=out_path,
               by_entity=by_entity, entity_total=len(findings))
    else:
        _audit(action="restore-file", user=_user(), session_id=sess.id if sess else None,
               ip=_ip(), kind=resp["kind"], filename=name, input_file=in_path,
               output_file=out_path, entity_total=len(findings))
    return jsonify(resp)


@app.post("/api/unmask-file")
def unmask_file_route():
    """Multipart: file + session_id. Restores tokens inside a file the LLM
    produced (Excel, Word, PowerPoint, Markdown/CSV/text) using the session
    vault. Returns the report and a download name; never changes the vault."""
    f = request.files.get("file")
    if not f or not f.filename:
        return jsonify({"error": "No file uploaded."}), 400
    name = os.path.basename(f.filename)
    ext = os.path.splitext(name)[1].lower()
    if ext not in RESTORE_EXTS:
        return jsonify({"error": f"Unsupported file type {ext!r} — supported: "
                                 f"{', '.join(RESTORE_EXTS)}."}), 415
    sess, err = _session_or_error(request.form.get("session_id"))
    if err:
        return err
    engine = get_engine(engine_options({}))
    run_id = uuid.uuid4().hex[:12]
    run_dir = os.path.join(RUNS, run_id)
    os.makedirs(run_dir)
    auth_mod.AUTH.runs.create(run_id, owner_id=_owner(), session_id=sess.id if sess else None)
    in_path = os.path.join(run_dir, "input" + ext)
    f.save(in_path)
    out_name = "restored" + ext
    t0 = time.time()
    try:
        with sessions.bind(engine, sess):
            report = unmask_file(engine, in_path, os.path.join(run_dir, out_name))
    except Exception as e:
        return jsonify({"error": f"Restore failed: {e}"}), 500
    resp = {"run_id": run_id, "session_id": sess.id, "filename": name,
            "elapsed_s": round(time.time() - t0, 2), **report,
            "downloads": {"output": out_name}}
    with open(os.path.join(run_dir, "result.json"), "w") as fh:
        json.dump(resp, fh)
    _audit(action="unmask-file", user=_user(), session_id=sess.id, ip=_ip(),
           kind=report.get("kind", "text"), filename=name,
           input_file=in_path, output_file=os.path.join(run_dir, out_name),
           entity_total=report.get("restored", 0), unresolved=report.get("unresolved", []))
    return jsonify(resp)


# ------------------------------------------------------------- audit API
# All /api/audit routes are gated (auditor role, or the admin key with auth off) in webui/auth.py.
@app.get("/api/audit")
def audit_list():
    def _float(name):
        v = request.args.get(name)
        try:
            return float(v) if v not in (None, "") else None
        except ValueError:
            return None
    try:
        limit = max(1, min(int(request.args.get("limit", 50)), 500))
    except ValueError:
        limit = 50
    try:
        offset = max(0, int(request.args.get("offset", 0)))
    except ValueError:
        offset = 0
    audit.maybe_sweep()
    return jsonify(audit.list(
        user=request.args.get("user") or None,
        action=request.args.get("action") or None,
        since=_float("since"), until=_float("until"),
        q=request.args.get("q") or None, limit=limit, offset=offset))


@app.get("/api/audit/stats")
def audit_stats():
    return jsonify(audit.stats())


@app.get("/api/audit/<rid>")
def audit_get(rid):
    rec = audit.get(rid)
    if not rec:
        return jsonify({"error": "Record not found."}), 404
    return jsonify(rec)


@app.get("/api/audit/<rid>/file/<role>")
def audit_file(rid, role):
    if role not in ("input", "output"):
        return jsonify({"error": "role must be input or output"}), 400
    path = audit.file_path(rid, role)
    if not path:
        return jsonify({"error": "File not found."}), 404
    rec = audit.get(rid) or {}
    stem = os.path.splitext(rec.get("filename") or role)[0]
    ext = os.path.splitext(path)[1]
    # For restore/unmask ops the input is the masked file and the output the
    # unmasked (restored) file; for mask ops it's the reverse.
    restore = rec.get("action") in ("unmask", "unmask-file", "restore-file")
    if role == "input":
        label = "masked" if restore else "original"
    else:
        label = "unmasked" if restore else "masked"
    return send_file(path, as_attachment=True, download_name=f"{stem}_{label}{ext}")


# ---- admin policy console (custom deny/allow/regex); admin-key gated ----
def _build_engine(options, overlay_data):
    """A one-off engine with a candidate overlay (not cached), sharing the
    already-loaded spaCy model. Used to preview unsaved rules."""
    model = options["nlp_model"]
    with _build_lock:
        nlp = _nlp_engines.get(model)
        if nlp is None:
            nlp = _nlp_engines[model] = build_nlp_engine(model)
    kw = dict(options, entities=list(options["entities"]) if options["entities"] else None)
    return FinancialPrivacyEngine(nlp_engine=nlp, overlay=overlay_data, **kw)


def _builtin_view():
    """Read-only summary of the built-in policy so an admin can see what is
    already covered before adding a custom rule."""
    pol = get_engine(engine_options({})).policy
    return {
        "locale": {"code": pol.code, "name": pol.name},
        "column_rules": [{"entity": e, "header": c.pattern} for (e, c, _d) in pol.column_rules],
        "profile_patterns": [{"entity": e, "regex": rx.pattern} for (e, rx, _v) in pol.profile_patterns],
        "disabled_entities": sorted(pol.disabled_entities),
        "locale_identifiers": sorted({i["entity"] for i in pol.identifiers}),
    }


@app.get("/api/policy")
def policy_get():
    _refresh_overlay()
    return jsonify({
        "overlay": _overlay,
        "builtins": _builtin_view(),
        "default_entity": overlay_mod.DEFAULT_ENTITY,
        "limits": {"min_deny_len": overlay_mod.MIN_DENY_LEN,
                   "max_regex_len": overlay_mod.MAX_REGEX_LEN},
    })


@app.put("/api/policy")
def policy_put():
    global _overlay, _overlay_fp, _overlay_rev, _overlay_checked
    data = request.get_json(silent=True) or {}
    try:
        norm, rev = _policy.save(data, user_id=_user())
    except overlay_mod.PolicyError as e:
        return jsonify({"error": str(e)}), 400
    with _build_lock:
        _overlay, _overlay_rev = norm, rev
        _overlay_fp = overlay_mod.fingerprint(norm)
        _overlay_checked = time.monotonic()
        _engines.clear()  # global change: drop every warm engine so it rebuilds
    counts = {"deny_terms": len(norm["deny_terms"]), "allow_terms": len(norm["allow_terms"]),
              "regex_rules": len(norm["regex_rules"])}
    _audit(action="policy-update", user=_user(), ip=_ip(), kind="policy",
           by_entity=counts, entity_total=sum(counts.values()),
           output_text=f"deny={counts['deny_terms']} allow={counts['allow_terms']} "
                       f"regex={counts['regex_rules']}")
    return jsonify({"ok": True, "overlay": norm})


@app.post("/api/policy/test")
def policy_test():
    data = request.get_json(silent=True) or {}
    text = (data.get("text") or "")[:MAX_TEXT_CHARS]
    if not text.strip():
        return jsonify({"error": "Provide some text to test."}), 400
    # Preview against the candidate overlay if supplied, else the live one.
    _refresh_overlay()
    candidate = data.get("overlay", _overlay)
    try:
        eng = _build_engine(engine_options({}), candidate)
    except overlay_mod.PolicyError as e:
        return jsonify({"error": str(e)}), 400
    masked, changed = eng.pseudonymize_text(text)
    # Count the tokens actually applied (grouped by entity), not raw findings:
    # overlapping detections that lost to a higher-priority rule never produce a
    # token, so counting findings would show phantom entities.
    by_entity = {}
    for tok in rules.TOKEN_RE.findall(masked):
        m = _re.match(r"TOK_(.+)_[0-9A-F]{8,}$", tok)
        if m:
            by_entity[m.group(1)] = by_entity.get(m.group(1), 0) + 1
    return jsonify({"masked": masked, "changed": changed,
                    "findings": eng._last_findings, "by_entity": by_entity})


def _run_file(run_id, fname):
    """Path of a run artefact the current principal may read, else None.
    Unknown runs and other principals' runs look the same (not found)."""
    safe = os.path.basename(fname)
    path = os.path.join(RUNS, os.path.basename(run_id), safe)
    if not os.path.isfile(path):
        return None
    owner = auth_mod.AUTH.runs.owner_of(run_id)
    if owner is auth_mod.MISSING:
        # A run dir with no row (made before the upgrade): ownerless.
        owner = None
    return path if owned(owner) else None


@app.get("/api/download/<run_id>/<path:fname>")
def download(run_id, fname):
    path = _run_file(run_id, fname)
    if not path:
        return jsonify({"error": "File not found"}), 404
    return send_file(path, as_attachment=True, download_name=os.path.basename(path))


@app.get("/api/text/<run_id>/<path:fname>")
def run_text(run_id, fname):
    """A run's text output as JSON (for copy-to-clipboard in the UI)."""
    path = _run_file(run_id, fname)
    if not path or not path.endswith((".md", ".txt")):
        return jsonify({"error": "File not found"}), 404
    with open(path, encoding="utf-8") as fh:
        return jsonify({"text": fh.read()})


if __name__ == "__main__":
    # Local default: loopback only. Containers/PaaS set HOST=0.0.0.0 and PORT.
    app.run(host=os.environ.get("HOST", "127.0.0.1"),
            port=int(os.environ.get("PORT", "5170")), debug=False)
